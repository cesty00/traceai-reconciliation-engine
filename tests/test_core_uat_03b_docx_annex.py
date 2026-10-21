"""Isolated synthetic compatibility tests. No production report changes."""
import importlib.util
from pathlib import Path
from docx import Document
import pytest

ROOT = Path(__file__).parents[1] / "prototypes/core_uat_03"
def load(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / (name + ".py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

adapter = load("operator_display")
annex = load("docx_annex")

def automatic():
    return {"case_id": "SYN-02", "product_code": "SYN-ITEM",
            "difference": -8.0, "residual": 0.0,
            "operator_status": "REVIEW",
            "wms_mutation_allowed": False, "wme_mutation_allowed": False,
            "evidence": [{"document_ref": "SYN-WMS-DOC",
                          "physical_row": 12, "source_hash": "sha-synthetic"}]}

def test_annex_separates_diagnostic_from_rationale(tmp_path):
    display = adapter.build_operator_display(automatic(), {
        "reason_code": "INVENTORY", "comment": "SYN count",
        "supporting_document": "SYN-PV", "reviewer": "TEST-USER",
        "validated_at": "2026-10-20T12:00:00Z", "scope": "SYN-WMS-DOC",
    })
    path = annex.render_validation_annex(display, tmp_path / "annex.docx")
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "CAUSE_IDENTIFIED" in text
    assert "OPERATOR_VALIDATED" in text
    assert "REVIEW" in text
    assert "NU autorizeaza ajustarea" in text
    assert "SYN-PV" in text

def test_incomplete_provenance_does_not_claim_cause(tmp_path):
    event = automatic()
    del event["evidence"][0]["physical_row"]
    display = adapter.build_operator_display(event)
    path = annex.render_validation_annex(display, tmp_path / "annex.docx")
    assert "Atribuire documentara: REVIEW" in "\n".join(p.text for p in Document(path).paragraphs)

def test_existing_report_is_never_overwritten(tmp_path):
    target = tmp_path / "case.docx"
    target.write_bytes(b"ORIGINAL")
    with pytest.raises(FileExistsError):
        annex.render_validation_annex(adapter.build_operator_display(automatic()), target)
    assert target.read_bytes() == b"ORIGINAL"

def test_missing_authority_rejected(tmp_path):
    display = adapter.build_operator_display(automatic())
    del display["wme_mutation_allowed"]
    with pytest.raises(ValueError):
        annex.render_validation_annex(display, tmp_path / "annex.docx")

def test_original_case_report_model_unchanged():
    from traceai_reconciliation.reports.report_model import CaseReport
    assert "operator_statement" not in CaseReport.__dataclass_fields__
    assert "initial_difference" in CaseReport.__dataclass_fields__
