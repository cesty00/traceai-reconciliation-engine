from pathlib import Path
from docx import Document
from traceai_reconciliation.reports.docx_report import _add_operator_causes
from traceai_reconciliation.reports.report_model import CaseReport


def test_causes_documentary_evidence_is_linked_by_exact_document(tmp_path: Path):
    report=CaseReport(
        case_id="CASE-T",product_code="T",product_name="TEST",period="2026",
        report_date="2026-10-15",initial_difference=0,explained_difference=0,
        unexplained_difference=0,unit="kg",verdict="REVIEW",evidence_level="B",
        cause_code="C",cause_description="Review",cause_status="NEDETERMINATA",
        executive_summary="Review",
        findings=[
            {"type":"SOURCE_COUNTS","description":"diagnostic","evidence":"SRC"},
            {"type":"RECEPTION_STAGING_INCLUDED_IN_STORAGE","description":"staging","evidence":"STOCK"},
            {"type":"STOCK_REPORT_OFFSETTING_ZONE_DIFFERENCES","description":"duplicate","evidence":"STOCK"},
        ],
        unmatched_operations=[{"system":"WME","document":"2613","operation_type":"PV","quantity":5.95,"reason":"review","date":"2026-08-29"}],
        operation_timeline=[
            {"system":"WME","document":"2613","date":"2026-08-29","warehouse":"DEPOZIT PRINCIPAL"},
            {"system":"WME","document":"26130","date":"2026-08-30","warehouse":"WRONG"},
        ],
    )
    doc=Document()
    _add_operator_causes(doc,report)
    output=tmp_path/"operator.docx"
    doc.save(output)
    rendered=Document(output)
    content=" ".join(p.text for p in rendered.paragraphs)
    tables=" ".join(c.text for t in rendered.tables for row in t.rows for c in row.cells)
    assert "SOURCE_COUNTS" not in content
    assert "STOCK_REPORT_OFFSETTING_ZONE_DIFFERENCES" not in content
    assert "RECEPTION_STAGING_INCLUDED_IN_STORAGE" in content
    assert "2613" in tables
    assert "DEPOZIT PRINCIPAL" in tables
    assert "WRONG" not in tables
