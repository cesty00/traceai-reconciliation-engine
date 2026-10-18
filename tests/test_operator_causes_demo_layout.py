from pathlib import Path
from docx import Document
from traceai_reconciliation.reports.report_model import CaseReport
from traceai_reconciliation.reports.operator_causes_report import generate_operator_causes_docx


def test_operator_report_follows_demo_structure(tmp_path: Path):
    report = CaseReport(
        case_id="CASE-TEST", product_code="SKU", product_name="TEST",
        period="01.01.2026 - 11.10.2026", report_date="2026-10-15",
        initial_difference=0, explained_difference=0, unexplained_difference=0,
        unit="kg", verdict="REVIEW", evidence_level="B", cause_code="C",
        cause_description="Review", cause_status="NEDETERMINATA",
        executive_summary="Test",
        findings=[{"type":"RECEPTION_STAGING_INCLUDED_IN_STORAGE"}],
        operation_timeline=[
            dict(system="WMS",document="SFA24774",operation_type="Livrare",quantity=-2700,date="2026-10-01 12:23:48",warehouse="Depozit Principal"),
            dict(system="WMS",document="SFA24774-1",operation_type="Receptie",quantity=300,date="2026-10-02 14:29:24",warehouse="Depozit Principal"),
            dict(system="WME",document="24774",operation_type="AE",quantity=-2400,date="2026-10-01",warehouse="DEPOZIT PRINCIPAL"),
        ],
        unmatched_operations=[dict(system="WME",document="2613",operation_type="PV",quantity=5.95,reason="PV fara corespondent",date="2026-08-29")]
    )
    path = generate_operator_causes_docx(
        report, tmp_path/"case.docx",
        details={"stock_report_zone_matching":{"rows":[
            dict(report_warehouse="Depozitare",quantity_wms=23380,quantity_erp=24280,report_difference=-900),
            dict(report_warehouse="Receptie",quantity_wms=900,quantity_erp=0,report_difference=900),
        ]}}
    )
    rendered=Document(path)
    prose=" ".join(p.text for p in rendered.paragraphs)
    tables=" ".join(c.text for t in rendered.tables for row in t.rows for c in row.cells)
    for label in ["RAPORT OPERATOR", "Rezumat stoc", "Registru CAUSES",
                  "Recepție nerelocată", "Livrare 24774", "Inventare și corecții",
                  "LOCALIZARE RAPIDĂ", "Decizia operatorului"]:
        assert label in prose
    for term in ["SFA24774-1", "2613", "23,380.000", "900.000"]:
        assert term in tables
    assert "Anexa A - Jurnal cronologic" not in prose

def test_unequal_stock_has_correct_warning(tmp_path: Path):
    report = CaseReport(
        case_id="CASE-40024", product_code="40024", product_name="Etichete",
        period="2026", report_date="2026-10-15", initial_difference=-8620.2,
        explained_difference=0, unexplained_difference=-8620.2,
        unit="BUCATA", verdict="REVIEW", evidence_level="B",
        cause_code="REVIEW", cause_description="Pending",
        cause_status="NEDETERMINATA", executive_summary="Review"
    )
    target = generate_operator_causes_docx(report,tmp_path/"case.docx",details={
        "stock_report_zone_matching":{"rows":[dict(report_warehouse="SECTIE PRODUCTIE",
            quantity_wms=18044.3,quantity_erp=26664.5,report_difference=-8620.2)]}})
    prose=" ".join(p.text for p in Document(target).paragraphs)
    assert "Diferență RAPORTATĂ, cu origine încă de verificat" in prose
    assert "Diferență de stoc REALĂ" not in prose
    assert "nu se poate confirma o diferență fizică" in prose
    assert "Un sold egal" not in prose
