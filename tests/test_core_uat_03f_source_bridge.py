"""CORE-UAT-03F synthetic bridge tests."""
import hashlib, importlib.util
from pathlib import Path
from openpyxl import Workbook

path=Path(__file__).parents[1]/"prototypes/core_uat_03/source_bridge.py"
spec=importlib.util.spec_from_file_location("bridge",path)
bridge=importlib.util.module_from_spec(spec);spec.loader.exec_module(bridge)

def test_wms_bridge_exact_source(tmp_path):
    p=tmp_path/"synthetic.csv";p.write_bytes(b"Document;Cod articol;Cantitate\nSYN-01;SYN-P;10\n")
    result=bridge.verify_evidence_pointer({
        "source_type":"WMS_CSV","source_path":str(p),
        "source_sha256":hashlib.sha256(p.read_bytes()).hexdigest(),
        "physical_row":2,"document_ref":"SYN-01","product_ref":"SYN-P",
        "expected_fields":{"Document":"SYN-01","Cod articol":"SYN-P","Cantitate":"10"}})
    assert result["provenance_status"]=="SOURCE_ROW_VERIFIED"
    assert result["economic_match_authorized"] is False
    assert result["adjustment_authorized"] is False

def test_wme_bridge_exact_xml_row(tmp_path):
    p=tmp_path/"synthetic.xlsx";wb=Workbook();ws=wb.active
    ws["C5"]="SYN-DOC";ws["K5"]="SYN-P";wb.save(p)
    result=bridge.verify_evidence_pointer({
        "source_type":"WME_XLSX","source_path":str(p),
        "source_sha256":hashlib.sha256(p.read_bytes()).hexdigest(),
        "worksheet_path":"xl/worksheets/sheet1.xml","physical_row":5,
        "document_ref":"SYN-DOC","product_ref":"SYN-P",
        "document_cell":"C5","product_cell":"K5",
        "expected_cells":{"C5":"SYN-DOC","K5":"SYN-P"}})
    assert result["provenance_status"]=="SOURCE_ROW_VERIFIED"
    assert result["operator_validated"] is False

def test_missing_identity_fails_closed():
    result=bridge.verify_evidence_pointer({"source_type":"WMS_CSV"})
    assert result["provenance_status"]=="REVIEW"

def test_unsupported_source_fails_closed():
    assert bridge.verify_evidence_pointer({"source_type":"UNKNOWN"})["provenance_status"]=="REVIEW"

def test_conflicting_identity_fails_closed():
    result=bridge.verify_evidence_pointer({"source_type":"WMS_CSV",
        "document_ref":"ONE","product_ref":"P",
        "expected_fields":{"Document":"TWO","Cod articol":"P"}})
    assert result["reason"]=="BUSINESS_IDENTITY_CONFLICT"
