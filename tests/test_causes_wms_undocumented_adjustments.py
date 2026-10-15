from pathlib import Path
from docx import Document
from traceai_reconciliation.reports.report_model import CaseReport
from traceai_reconciliation.reports.operator_causes_report import generate_operator_causes_docx

def test_40024_documentless_adjustments_preserve_signed_groups(tmp_path: Path):
    report=CaseReport(
        case_id="CASE-40024", product_code="40024", product_name="ETICHETE",
        period="2026", report_date="2026-10-15", initial_difference=-8620.2,
        explained_difference=0, unexplained_difference=-8620.2, unit="BUCATA",
        verdict="REVIEW", evidence_level="B", cause_code="REVIEW",
        cause_description="PENDING", cause_status="NEDETERMINATA",
        executive_summary="Review",
        operation_timeline=[
            dict(system="WMS",document="-",operation_type="Ajustare negativa",date="2026-08-12 17:19:22",quantity=-8620,source_row=275716,source_metadata={"lot":"260709"}),
            dict(system="WMS",document="-",operation_type="Ajustare negativa",date="2026-08-12 17:20:07",quantity=-0.2,source_row=275715,source_metadata={"lot":"260709"}),
            dict(system="WMS",document="-",operation_type="Ajustare pozitiva",date="2026-10-06 11:51:52",quantity=8620.2,source_row=275720,source_metadata={"lot":"260827"}),
            dict(system="WMS",document="-",operation_type="Ajustare negativa",date="2026-10-06 15:07:21",quantity=-8620.2,source_row=275714,source_metadata={"lot":"260827"}),
            dict(system="WMS",document="0023556_AE",operation_type="Ajustare negativa",date="2026-01-05",quantity=-1008,source_row=282154),
        ],
    )
    target=generate_operator_causes_docx(report,tmp_path/"case.docx")
    d=Document(target)
    all_text=" ".join(p.text for p in d.paragraphs)+" "+" ".join(c.text for tb in d.tables for row in tb.rows for c in row.cells)
    assert "Ajustări WMS fără document" in all_text
    for value in ["275716","275715","275720","275714","260709","260827"]:
        assert value in all_text
    assert "net -8,620.200" in all_text
    assert "net +0.000" in all_text
    assert "282154" not in all_text
    assert "Stadiu operațional: REVIEW" in all_text
