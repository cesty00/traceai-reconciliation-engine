"""GOV-002AA adversarial, test-only characterization of read-only provenance adapter.

COMPLETE means source *candidate set*, NOT lot-level/warehouse-level pairing
or confirmed economic equivalence. No CORE edits and no forced MATCH.
"""
from openpyxl import load_workbook

from test_gov_002w_provenance_probe import _sources
from traceai_reconciliation.investigation.provenance_adapter import resolve_supplier_receipt_provenance


def _call(data):
    wms, wme, production, index, *_ = data
    return resolve_supplier_receipt_provenance(
        index, wms_path=wms, wme_path=wme, production_path=production,
        product_code="SYN-W", product_name="Synthetic W", document="F88")


def test_gov_002aa_different_distribution_only_candidate_set(tmp_path):
    data = _sources(tmp_path)
    wme = data[1]
    wb = load_workbook(wme)
    wb.active["E3"] = 7
    wb.active["E4"] = 3
    wb.save(wme)
    data[3].ensure_index(data[0], data[1], data[2])
    outcome = _call(data)
    # Existing adapter accepts total 10 despite different 6/4 vs 7/3 split.
    # No assertion of per-line equivalence or economic ownership is permissible.
    assert outcome.status == "COMPLETE"
    assert len(outcome.source_lines) == 4


def test_gov_002aa_multiple_warehouses_only_candidate_set(tmp_path):
    data = _sources(tmp_path)
    wme = data[1]
    wb = load_workbook(wme)
    wb.active["N4"] = "OTHER WAREHOUSE"
    wb.save(wme)
    data[3].ensure_index(data[0], data[1], data[2])
    outcome = _call(data)
    # Existing adapter does not compare WMS/WME warehouse granularity.
    assert outcome.status == "COMPLETE"
    assert len(outcome.source_lines) == 4


def test_gov_002aa_unsupported_wme_operation_stays_unavailable(tmp_path):
    data = _sources(tmp_path)
    wb = load_workbook(data[1])
    wb.active["B4"] = "NT"
    wb.save(data[1])
    data[3].ensure_index(data[0], data[1], data[2])
    assert _call(data).status == "UNAVAILABLE"


def test_gov_002aa_changed_source_on_final_freshness_check_fail_closed(tmp_path, monkeypatch):
    data = _sources(tmp_path)
    index = data[3]
    original = index.is_current
    checks = {"count": 0}

    def is_current_with_mutation(*args, **kwargs):
        checks["count"] += 1
        if checks["count"] == 2:
            wms = data[0]
            wms.write_text(wms.read_text(encoding="utf-8").replace("6.0", "7.0"),
                           encoding="utf-8")
        return original(*args, **kwargs)

    monkeypatch.setattr(index, "is_current", is_current_with_mutation)
    result = _call(data)
    assert checks["count"] >= 2
    assert result.status == "UNAVAILABLE"
    assert result.reason == "STALE_SOURCE_INDEX"
    assert not result.source_lines
