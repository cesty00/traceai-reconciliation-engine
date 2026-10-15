from pathlib import Path
from docx import Document
from traceai_reconciliation.reports.docx_report import _add_operator_causes
from traceai_reconciliation.reports.report_model import CaseReport


def test_delivery_return_evidence_panel_is_document_specific(tmp_path: Path):
    report = CaseReport(
        case_id="TEST", product_code="X", product_name="X", period="2026",
        report_date="2026-10-15", initial_difference=0,
        explained_difference=0, unexplained_difference=0, unit="kg",
        verdict="REVIEW", evidence_level="B", cause_code="C",
        cause_description="Review", cause_status="NEDETERMINATA",
        executive_summary="Review",
        operation_timeline=[
            {"system":"WMS","document":"SFA24774","operation_type":"Livrare","date":"2026-10-01 12:23:48","quantity":-2700,"warehouse":"Depozit Principal"},
            {"system":"WMS","document":"SFA24774-1","operation_type":"Receptie","date":"2026-10-02 14:29:24","quantity":300,"warehouse":"Depozit Principal"},
            {"system":"WME","document":"24774","operation_type":"AE","date":"2026-10-01","quantity":-2400,"warehouse":"DEPOZIT PRINCIPAL"},
            {"system":"WME","document":"247740","operation_type":"AE","date":"2026-10-01","quantity":-10,"warehouse":"OTHER"}
        ])
    document = Document()
    _add_operator_causes(document, report)
    output = tmp_path / "test.docx"
    document.save(output)
    loaded = Document(output)
    prose = " ".join(p.text for p in loaded.paragraphs)
    cells = " ".join(cell.text for table in loaded.tables for row in table.rows for cell in row.cells)
    assert "Document 24774" in prose
    assert "WMS net -2400.000 vs AE WME -2400.000" in prose
    assert "SFA24774-1" in cells
    assert "247740" not in cells
