"""GOV-002Z: source-backed isolated read-only adapter integration tests."""
from openpyxl import load_workbook

from test_gov_002w_provenance_probe import _sources
from traceai_reconciliation.investigation.provenance_adapter import (
    resolve_supplier_receipt_provenance,
)


def _resolve(data):
    wms, wme, production, index, *_ = data
    return resolve_supplier_receipt_provenance(
        index, wms_path=wms, wme_path=wme, production_path=production,
        product_code="SYN-W", product_name="Synthetic W", document="F88",
    )


def test_gov_002z_exact_receipt_has_four_physical_rows_and_hashes(tmp_path):
    data = _sources(tmp_path)
    result = _resolve(data)
    assert result.status == "COMPLETE"
    assert {(x.system, x.physical_row_number) for x in result.source_lines} == {
        ("WMS", 2), ("WMS", 3), ("WME", 3), ("WME", 4)}
    assert all(len(x.source_sha256) == 64 for x in result.source_lines)


def test_gov_002z_stale_wms_fails_closed(tmp_path):
    data = _sources(tmp_path)
    wms = data[0]
    wms.write_text(wms.read_text(encoding="utf-8").replace("6.0", "7.0"),
                   encoding="utf-8")
    result = _resolve(data)
    assert result.status == "UNAVAILABLE"
    assert not result.source_lines


def test_gov_002z_stale_wme_fails_closed(tmp_path):
    data = _sources(tmp_path)
    wme = data[1]
    book = load_workbook(wme)
    book.active["E3"] = 7
    book.save(wme)
    assert _resolve(data).status == "UNAVAILABLE"


def test_gov_002z_missing_document_fails_closed(tmp_path):
    data = _sources(tmp_path)
    wms, wme, production, index, *_ = data
    result = resolve_supplier_receipt_provenance(
        index, wms_path=wms, wme_path=wme, production_path=production,
        product_code="SYN-W", product_name="Synthetic W", document="NOT-FOUND"
    )
    assert result.status == "UNAVAILABLE"
    assert result.reason == "SOURCE_SET_INCOMPLETE"


def test_gov_002z_reused_document_other_product_fails_closed(tmp_path):
    data = _sources(tmp_path)
    wms, wme, production, index, *_ = data
    book = load_workbook(wme)
    book.active.append([3, "FE", "F88", "02.01.2026", 1, 0, 1, 0,
                        "345", "SYN", "OTHER PRODUCT", "", 222, "DP", "BUC", "TEST", 3])
    book.save(wme)
    assert index.ensure_index(wms, wme, production)
    result = _resolve(data)
    assert result.status == "UNAVAILABLE"
    assert result.reason == "DOCUMENT_REUSED_ACROSS_PRODUCTS"


def test_gov_002z_unsupported_wme_document_type_fails_closed(tmp_path):
    data = _sources(tmp_path)
    wms, wme, production, index, *_ = data
    book = load_workbook(wme)
    book.active["B3"] = "F"
    book.save(wme)
    assert index.ensure_index(wms, wme, production)
    result = _resolve(data)
    assert result.status == "UNAVAILABLE"
    assert result.reason == "UNSUPPORTED_DOCUMENT_FAMILY"
