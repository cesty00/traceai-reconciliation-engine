from pathlib import Path
from docx import Document
from traceai_reconciliation.reports.report_model import CaseReport
from traceai_reconciliation.reports.operator_causes_report import generate_operator_causes_docx

def test_operator_pv_and_mose_rows_are_from_exact_source(tmp_path: Path):
    report = CaseReport(
        case_id="CASE-X", product_code="X", product_name="X", period="2026",
        report_date="2026-10-15", initial_difference=0, explained_difference=0,
        unexplained_difference=0, unit="kg", verdict="REVIEW",
        evidence_level="B", cause_code="REVIEW", cause_description="Review",
        cause_status="NEDETERMINATA", executive_summary="Review",
        unmatched_operations=[
            dict(system="WME", document="2446", operation_type="PV", quantity=0.9, reason="PV"),
            dict(system="WME", document="92766", operation_type="F", quantity=20.0, reason="MOSE"),
        ],
        operation_timeline=[
            dict(system="WME", document="2446", operation_type="PV", quantity=-0.9, source_row=111978, date="2026-06-04", warehouse="DEPOZIT PRINCIPAL"),
            dict(system="WME", document="92766", operation_type="F", quantity=-20.0, source_row=112037, date="2026-05-06", warehouse="MOSE"),
        ])
    out = generate_operator_causes_docx(report, tmp_path / "case.docx")
    doc = Document(out)
    content = " ".join(c.text for table in doc.tables for row in table.rows for c in row.cells)
    assert "111978" in content
    assert "112037" in content
    assert "-20.0" in content

def test_unresolved_pv_cannot_borrow_arbitrary_source_row(tmp_path: Path):
    report = CaseReport(
        case_id="CASE-Y", product_code="X", product_name="X", period="2026",
        report_date="2026-10-15", initial_difference=0, explained_difference=0,
        unexplained_difference=0, unit="kg", verdict="REVIEW",
        evidence_level="B", cause_code="REVIEW", cause_description="Review",
        cause_status="NEDETERMINATA", executive_summary="Review",
        unmatched_operations=[dict(system="WMS", document="-", operation_type="INV-MIN", quantity=-5.95, reason="No documented ID")],
        operation_timeline=[dict(system="WMS", document="-", operation_type="INV-MIN", quantity=-5.95, source_row=1234, date="2026-08-05")])
    out = generate_operator_causes_docx(report, tmp_path / "case.docx")
    doc = Document(out)
    content = " ".join(c.text for table in doc.tables for row in table.rows for c in row.cells)
    assert "Nedisponibil în cronologia cazului" in content
    assert "1234" not in content
