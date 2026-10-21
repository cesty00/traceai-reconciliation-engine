"""Synthetic WME XLSX physical XML row assertions, never real company exports."""
import hashlib
import importlib.util
from pathlib import Path
from openpyxl import Workbook

path = Path(__file__).parents[1] / "prototypes/core_uat_03/wme_xlsx_evidence.py"
spec = importlib.util.spec_from_file_location("uat03_wme_xlsx", path)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)

def fixture(tmp_path):
    p = tmp_path / "synthetic.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "LineNo"; ws["C1"] = "Document"
    ws["A5"] = "99"; ws["C5"] = "SYN-DOC"; ws["K5"] = "SYN-ITEM"
    wb.save(p)
    return p, hashlib.sha256(p.read_bytes()).hexdigest()

def check(p, sha, row=5, sheet="xl/worksheets/sheet1.xml", cells=None):
    return module.verify_wme_xlsx_evidence(p, expected_sha256=sha,
        worksheet_path=sheet, physical_row=row,
        expected_cells=cells or {"C5": "SYN-DOC", "K5": "SYN-ITEM"})

def test_actual_xml_row_not_column_a(tmp_path):
    p,sha=fixture(tmp_path)
    assert check(p,sha)["verified"] is True
    assert check(p,sha,row=99)["verified"] is False

def test_wrong_document_fails(tmp_path):
    p,sha=fixture(tmp_path)
    assert check(p,sha,cells={"C5":"OTHER"})["reason"]=="DOCUMENT_OR_CELL_MISMATCH"

def test_changed_file_invalidates_evidence(tmp_path):
    p,sha=fixture(tmp_path)
    with p.open("ab") as f: f.write(b"change")
    assert check(p,sha)["reason"]=="SOURCE_SHA_MISMATCH"

def test_missing_worksheet_fails(tmp_path):
    p,sha=fixture(tmp_path)
    assert check(p,sha,sheet="xl/worksheets/missing.xml")["reason"]=="WORKSHEET_NOT_FOUND"

def test_missing_row_fails(tmp_path):
    p,sha=fixture(tmp_path)
    assert check(p,sha,row=6)["reason"]=="AMBIGUOUS_OR_MISSING_ROW"

def test_invalid_pointer_fails(tmp_path):
    p,sha=fixture(tmp_path)
    assert check(p,sha,row=0)["reason"]=="INVALID_POINTER"
