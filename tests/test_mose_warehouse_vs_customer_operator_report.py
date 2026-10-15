from pathlib import Path
from docx import Document
from traceai_reconciliation.reports.report_model import CaseReport
from traceai_reconciliation.reports.operator_causes_report import generate_operator_causes_docx


def test_mose_warehouse_not_customer_name(tmp_path: Path):
    report = CaseReport(
        case_id="MOSE-TEST", product_code="SKU", product_name="Test", period="2026",
        report_date="2026-10-15", initial_difference=0,
        explained_difference=0, unexplained_difference=0, unit="kg",
        verdict="REVIEW", evidence_level="B", cause_code="C",
        cause_description="Review", cause_status="NEDETERMINATA",
        executive_summary="Review",
        operation_timeline=[
            dict(system="WME", document="92766", operation_type="F",
                 quantity=-20, date="2026-05-06", warehouse="MOSE",
                 counterpart="MONTELUX"),
            dict(system="WME", document="92767", operation_type="F",
                 quantity=-20, date="2026-05-06", warehouse="DEPOZIT PRINCIPAL",
                 counterpart="MOSE Business Distribution"),
        ],
        unmatched_operations=[
            dict(system="WME", document="92766", operation_type="F", quantity=20, reason="MOSE"),
            dict(system="WME", document="92767", operation_type="F", quantity=20,
                 reason="MOSE Business Distribution client"),
        ],
    )
    doc = Document(generate_operator_causes_docx(
        report, tmp_path/"case.docx",
        details={"transfer_mose_matching":{"matches":1}},
    ))
    paragraphs = " ".join(p.text for p in doc.paragraphs)
    tables = " ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "REGULĂ MOSE CONFIRMATĂ" in paragraphs
    assert "TRANSFER_MOSE" in paragraphs
    assert "MOSE Business Distribution" in paragraphs
    assert "92766" in tables
    assert "92767" not in tables
    assert "-20" in tables
