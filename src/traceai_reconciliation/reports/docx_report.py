from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .report_model import CaseReport

VERDICTS = {"MATCH_DIRECT", "MATCH_NORMALIZARE", "MATCH_DUPA_NETARE", "REVIEW", "NOT_MATCH"}
EVIDENCE_LEVELS = {"A", "B", "C", "D"}
CAUSE_STATUSES = {"CONFIRMATA", "PROBABILA", "NEDETERMINATA"}
CASE_STATUSES = {"DESCHIS", "IN_ANALIZA", "ASTEPTARE_VALIDARE", "INCHIS"}
RECOMMENDATION_TYPES = {"IMEDIATA", "CORECTIVA", "PREVENTIVA", "AUTOMATIZARE_CONTROL"}

DOCX_ANNEX_ROW_LIMIT = 120


def _sample_rows_for_docx(
    rows: Sequence[Sequence[object]],
    limit: int = DOCX_ANNEX_ROW_LIMIT,
) -> tuple[list[Sequence[object]], int]:
    """Keep report annexes usable while the technical JSON retains all rows."""
    materialized = list(rows)
    total = len(materialized)
    if total <= limit:
        return materialized, 0

    head_count = limit // 2
    tail_count = limit - head_count
    selected = (
        materialized[:head_count]
        + materialized[-tail_count:]
    )
    return selected, total - len(selected)


def validate_case_report(report: CaseReport, tolerance: float = 0.001) -> list[str]:
    errors: list[str] = []
    required_text = {
        "case_id": report.case_id,
        "product_code": report.product_code,
        "product_name": report.product_name,
        "period": report.period,
        "report_date": report.report_date,
        "unit": report.unit,
        "cause_code": report.cause_code,
        "cause_description": report.cause_description,
        "executive_summary": report.executive_summary,
    }
    for field_name, value in required_text.items():
        if not str(value or "").strip():
            errors.append(f"Camp obligatoriu lipsa: {field_name}")

    if report.verdict not in VERDICTS:
        errors.append(f"Verdict invalid: {report.verdict}")
    if report.evidence_level not in EVIDENCE_LEVELS:
        errors.append(f"Nivel de proba invalid: {report.evidence_level}")
    if report.cause_status not in CAUSE_STATUSES:
        errors.append(f"Status cauza invalid: {report.cause_status}")
    if report.case_status not in CASE_STATUSES:
        errors.append(f"Status caz invalid: {report.case_status}")

    reconciliation_gap = report.initial_difference - (
        report.explained_difference + report.unexplained_difference
    )
    if abs(reconciliation_gap) > tolerance:
        errors.append(
            "Inconsistenta diferente: initial_difference trebuie sa fie egal cu "
            "explained_difference + unexplained_difference in toleranta configurata."
        )

    for idx, recommendation in enumerate(report.recommendations, start=1):
        if recommendation.action_type not in RECOMMENDATION_TYPES:
            errors.append(
                f"Tip recomandare invalid la pozitia {idx}: {recommendation.action_type}"
            )
        if not recommendation.recommendation.strip():
            errors.append(f"Recomandare goala la pozitia {idx}")

    return errors


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=80, start=80, bottom=80, end=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)


def _fmt_qty(value: float | int | None) -> str:
    if value is None:
        return "-"
    return f"{float(value):,.3f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _text(value) -> str:
    if value is None or value == "":
        return "-"
    return str(value)


def _configure_styles(doc: DocumentObject) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for style_name, size in (("Title", 18), ("Heading 1", 13), ("Heading 2", 11)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(31, 56, 100)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)


def _configure_page(doc: DocumentObject) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    header = section.header.paragraphs[0]
    header.text = "TraceAI Control | Reconciliere WMS ↔ WME"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if header.runs:
        header.runs[0].font.name = "Arial"
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    footer = section.footer.paragraphs[0]
    footer.text = (
        "Raport de verificare produs | Document generat automat - "
        "necesita validare umana unde este indicat"
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if footer.runs:
        footer.runs[0].font.name = "Arial"
        footer.runs[0].font.size = Pt(7.5)
        footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)


def _add_heading(doc: DocumentObject, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _add_key_value_table(doc: DocumentObject, rows: Sequence[tuple[str, str]], widths=(5.2, 11.8)):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for label, value in rows:
        row = table.add_row()
        _prevent_row_split(row)
        cells = row.cells
        cells[0].text = label
        cells[1].text = value
        cells[0].width = Cm(widths[0])
        cells[1].width = Cm(widths[1])
        _set_cell_shading(cells[0], "EAF0F8")
        for i, cell in enumerate(cells):
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    if i == 0:
                        run.font.bold = True
    return table


def _add_table(
    doc: DocumentObject,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    font_size: float = 8.0,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        _set_cell_shading(header_cells[i], "D9E2F3")
        _set_cell_margins(header_cells[i], top=55, bottom=55, start=55, end=55)
        for paragraph in header_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(font_size)
    _repeat_table_header(table.rows[0])
    _prevent_row_split(table.rows[0])

    for row in rows:
        table_row = table.add_row()
        _prevent_row_split(table_row)
        cells = table_row.cells
        for i, value in enumerate(row):
            cells[i].text = _text(value)
            _set_cell_margins(cells[i], top=45, bottom=45, start=45, end=45)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
    return table


def _add_bullets(doc: DocumentObject, items: Iterable[str]) -> None:
    count = 0
    for item in items:
        count += 1
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(str(item))
    if count == 0:
        doc.add_paragraph("Nu au fost înregistrate elemente pentru această secțiune.")


def generate_docx_report(report: CaseReport, output_path: str | Path) -> Path:
    errors = validate_case_report(report)
    if errors:
        raise ValueError("Raport invalid: " + " | ".join(errors))

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _configure_styles(doc)
    _configure_page(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("RAPORT DE VERIFICARE PRODUS WMS ↔ WME")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor(31, 56, 100)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(f"{report.case_id} | {report.product_code}")
    subtitle_run.bold = True
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(11)

    _add_key_value_table(
        doc,
        [
            ("Cod produs", report.product_code),
            ("Denumire produs", report.product_name),
            ("Perioadă verificată", report.period),
            ("Data raportului", report.report_date),
            ("Status caz", report.case_status),
        ],
    )

    _add_heading(doc, "1. Rezumat executiv")
    doc.add_paragraph(report.executive_summary)
    _add_key_value_table(
        doc,
        [
            ("Diferență inițială", f"{_fmt_qty(report.initial_difference)} {report.unit}"),
            ("Diferență explicată", f"{_fmt_qty(report.explained_difference)} {report.unit}"),
            ("Diferență neexplicată", f"{_fmt_qty(report.unexplained_difference)} {report.unit}"),
            ("Verdict", report.verdict),
            ("Cauză", f"{report.cause_code} - {report.cause_description}"),
            ("Status cauză", report.cause_status),
            ("Nivel de probă", report.evidence_level),
            ("Necesită intervenție umană", "DA" if report.human_intervention_required else "NU"),
        ],
    )

    _add_heading(doc, "2. Obiectul și limitele verificării")
    doc.add_paragraph(
        "Verificarea urmărește concordanța scriptică și operațională dintre WMS și WME pentru "
        "produsul indicat, în perioada declarată. Raportul nu autorizează automat ajustări de "
        "stoc, reoperări de documente sau regularizări contabile; aceste acțiuni necesită "
        "validare umană conform responsabilităților interne."
    )

    _add_heading(doc, "3. Situația de pornire")
    _add_key_value_table(
        doc,
        [
            ("Diferență raportată", f"{_fmt_qty(report.initial_difference)} {report.unit}"),
            ("Unitate de măsură", report.unit),
            ("CASE_ID", report.case_id),
            ("Produs", f"{report.product_code} - {report.product_name}"),
        ],
    )

    _add_heading(doc, "4. Surse și dovezi verificate")
    if report.sources_used:
        rows = [
            (
                source.get("source_id", "-"),
                source.get("name", "-"),
                source.get("period", "-"),
                source.get("role", "-"),
                source.get("status", "VERIFICAT"),
            )
            for source in report.sources_used
        ]
        _add_table(doc, ["Sursă", "Fișier / registru", "Perioadă", "Rol", "Status"], rows, 7.5)
    else:
        doc.add_paragraph("Nu au fost înregistrate surse în rezultatul tehnic al cazului.")

    _add_heading(doc, "5. Metoda de verificare aplicată")
    _add_bullets(
        doc,
        [
            "Verificarea stocului inițial și a soldului raportat.",
            "Reconstrucția cronologică a mișcărilor relevante WMS și WME.",
            "Verificarea operațiunilor de producție, transfer, retur, ajustare și integrare, după caz.",
            "Aplicarea regulilor validate de normalizare și netare.",
            "Separarea verdictului de cauza identificată și atribuirea nivelului de probă.",
        ],
    )

    _add_heading(doc, "6. Rezultatul reconcilierii")
    _add_table(
        doc,
        ["Indicator", "Cantitate", "Interpretare"],
        [
            ("Diferență inițială", f"{_fmt_qty(report.initial_difference)} {report.unit}", "Semnalul de pornire al cazului"),
            ("Diferență explicată", f"{_fmt_qty(report.explained_difference)} {report.unit}", "Acoperită prin dovezi/reguli validate"),
            (
                "Diferență neexplicată",
                f"{_fmt_qty(report.unexplained_difference)} {report.unit}",
                "Rămâne deschisă pentru investigație dacă este diferită de zero",
            ),
        ],
        8.0,
    )

    _add_heading(doc, "7. Constatări")
    if report.findings:
        rows = [
            (
                idx,
                finding.get("type", "CONSTATARE"),
                finding.get("description", "-"),
                finding.get("evidence", "-"),
            )
            for idx, finding in enumerate(report.findings, start=1)
        ]
        _add_table(doc, ["Nr.", "Tip", "Constatare", "Dovadă / referință"], rows, 7.8)
    else:
        doc.add_paragraph("Nu au fost înregistrate constatări suplimentare.")

    _add_heading(doc, "8. Analiza cauzei")
    _add_key_value_table(
        doc,
        [
            ("Cod cauză", report.cause_code),
            ("Descriere", report.cause_description),
            ("Status cauză", report.cause_status),
            ("Nivel de probă", report.evidence_level),
        ],
    )

    _add_heading(doc, "9. Impactul constatat")
    if report.unexplained_difference == 0:
        doc.add_paragraph(
            "Diferența inițială este explicată integral în limitele regulilor și dovezilor "
            "utilizate. Acest fapt nu implică automat efectuarea unei corecții în sistem."
        )
    else:
        doc.add_paragraph(
            f"Rămâne o diferență neexplicată de {_fmt_qty(report.unexplained_difference)} "
            f"{report.unit}. Cazul necesită continuarea investigației sau validare umană "
            "înaintea oricărei corecții."
        )

    # G14 operator action overview: presentation-only, no balance effects.
    pv_actions = [
        item for item in report.unmatched_operations
        if item.get("operation_type") == "COR_STOC_PV_DOCUMENTARY_CANDIDATE"
        and item.get("status") == "COR_STOC_PV_SOURCE_REVIEW"
    ]
    if pv_actions:
        _add_heading(doc, "9A. Acțiuni necesare operatorului — documente PV")
        doc.add_paragraph(
            "Candidați documentari pentru verificare umană. Aceleași operațiuni "
            "pot apărea și în anexele cu reziduuri WMS/WME; rândurile de mai jos "
            "NU reprezintă diferențe suplimentare de stoc, NU se însumează și "
            "NU autorizează corecții sau neutralizări automate."
        )
        def _pv_action(item):
            details = str(item.get("reason", ""))
            if "gestiune DE VERIFICAT" in details:
                priority, action = "P1", "Verificați gestiunea documentului PV și transferurile dintre gestiuni."
            elif "cantități WMS=" in details and "diferență documentară=0.0;" not in details:
                priority, action = "P1", "Verificați rândurile WMS neasociate explicit și documentul PV."
            else:
                priority, action = "P2", "Confirmați legătura documentară și validarea operațională."
            return (priority, str(item.get("document", "-")), action, details)
        # Sort using the same computed priority shown in the table; otherwise
        # a P1 quantity discrepancy can appear after P2 documentary checks.
        pv_actions.sort(key=lambda item: (
            0 if _pv_action(item)[0] == "P1" else 1,
            str(item.get("document", ""))
        ))
        _add_table(
            doc,
            ["Prioritate", "PV", "Verificarea operatorului", "Dovezi / rânduri originale"],
            [_pv_action(item) for item in pv_actions],
            7.2,
        )

    _add_heading(doc, "10. Recomandări de rezolvare")
    if report.recommendations:
        rows = [
            (
                recommendation.priority,
                recommendation.action_type,
                recommendation.recommendation,
                recommendation.proposed_owner or "-",
                recommendation.precondition or "-",
            )
            for recommendation in report.recommendations
        ]
        _add_table(
            doc,
            ["Prioritate", "Tip", "Recomandare", "Responsabil propus", "Precondiție"],
            rows,
            7.3,
        )
    else:
        doc.add_paragraph("Nu au fost generate recomandări automate pentru acest caz.")

    note = doc.add_paragraph()
    note_run = note.add_run("Notă de control: ")
    note_run.bold = True
    note.add_run(
        "recomandările sunt propuneri de remediere și prevenție. Ele nu reprezintă aprobarea "
        "unei ajustări WMS/WME sau a unei regularizări contabile."
    )

    _add_heading(doc, "11. Verdict și închidere caz")
    _add_key_value_table(
        doc,
        [
            ("Verdict final", report.verdict),
            ("Cauză", report.cause_code),
            ("Nivel de probă", report.evidence_level),
            ("Status caz", report.case_status),
            ("Necesită intervenție umană", "DA" if report.human_intervention_required else "NU"),
        ],
    )

    _add_heading(doc, "12. Validare umană")
    _add_table(
        doc,
        ["Rol", "Nume", "Decizie / observație", "Data"],
        [
            ("Verificator", "", "", ""),
            ("Responsabil operațional", "", "", ""),
            ("Aprobator, dacă este necesar", "", "", ""),
        ],
        8.0,
    )

    section = doc.add_section()
    _set_landscape(section)

    _add_heading(doc, "Anexa A - Jurnal cronologic al operațiunilor relevante")
    timeline_rows = [
        (
            operation.get("date", "-"),
            operation.get("system", "-"),
            operation.get("document", "-"),
            operation.get("operation_type", "-"),
            operation.get("quantity", "-"),
            operation.get("warehouse", "-"),
            operation.get("counterpart", "-"),
            operation.get("verdict", "-"),
            operation.get("note", "-"),
        )
        for operation in report.operation_timeline
    ]
    if timeline_rows:
        timeline_rows_docx, timeline_omitted = _sample_rows_for_docx(
            timeline_rows
        )
        if timeline_omitted:
            doc.add_paragraph(
                f"Anexa Word afișează {len(timeline_rows_docx)} din "
                f"{len(timeline_rows)} operațiuni (primele și ultimele din "
                f"cronologie). {timeline_omitted} rânduri sunt omise din "
                "DOCX pentru lizibilitate; jurnalul complet rămâne în "
                "rezultatul tehnic JSON al execuției."
            )
        _add_table(
            doc,
            ["Data", "Sistem", "Document", "Tip operațiune", "Cantitate", "Gestiune", "Corespondent", "Verdict", "Observație"],
            timeline_rows_docx,
            6.6,
        )
    else:
        doc.add_paragraph("Nu există operațiuni cronologice anexate.")

    _add_heading(doc, "Anexa B - Operațiuni fără corespondent / diferențe punctuale")
    unmatched_rows = [
        (
            item.get("date", "-"),
            item.get("system", "-"),
            item.get("document", "-"),
            item.get("operation_type", "-"),
            item.get("quantity", "-"),
            item.get("reason", "-"),
            item.get("status", "-"),
        )
        for item in report.unmatched_operations
    ]
    if unmatched_rows:
        unmatched_rows_docx, unmatched_omitted = _sample_rows_for_docx(
            unmatched_rows
        )
        if unmatched_omitted:
            doc.add_paragraph(
                f"Anexa Word afișează {len(unmatched_rows_docx)} din "
                f"{len(unmatched_rows)} operațiuni fără corespondent / "
                f"diferențe punctuale. {unmatched_omitted} rânduri sunt "
                "omise din DOCX pentru lizibilitate; lista completă rămâne "
                "în rezultatul tehnic JSON al execuției."
            )
        _add_table(
            doc,
            ["Data", "Sistem", "Document", "Tip", "Cantitate", "Motiv", "Status"],
            unmatched_rows_docx,
            7.0,
        )
    else:
        doc.add_paragraph("Nu există operațiuni fără corespondent anexate.")

    _add_heading(doc, "Anexa C - Trasabilitatea concluziei")
    trace_rows = [
        (
            item.get("step", "-"),
            item.get("source", "-"),
            item.get("key", "-"),
            item.get("rule", "-"),
            item.get("result", "-"),
            item.get("evidence_level", "-"),
        )
        for item in report.conclusion_traceability
    ]
    if trace_rows:
        _add_table(
            doc,
            ["Pas", "Sursă", "Cheie", "Regulă", "Rezultat", "Nivel probă"],
            trace_rows,
            7.0,
        )
    else:
        doc.add_paragraph("Trasabilitatea concluziei nu a fost încă populată.")

    _add_heading(doc, "Anexa D - Parametri tehnici ai execuției")
    hashes = "; ".join(f"{key}={value}" for key, value in sorted(report.source_hashes.items())) or "-"
    _add_key_value_table(
        doc,
        [
            ("Execution ID", report.execution_id or "-"),
            ("Versiune motor", report.engine_version or "-"),
            ("Commit GitHub", report.engine_commit_sha or "-"),
            ("Versiune reguli", report.rules_version or "-"),
            ("Hash reguli", report.rules_hash or "-"),
            ("Hash-uri surse", hashes),
        ],
        widths=(5.5, 20.0),
    )

    if report.technical_warnings:
        doc.add_paragraph("Avertismente tehnice:")
        _add_bullets(doc, report.technical_warnings)

    doc.save(output_path)
    return output_path
