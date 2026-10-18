"""Concise operator-first CAUSES report; full forensic report remains separate.

No hardcoded product IDs or assumed evidence. Original source rows are kept in JSON.
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from traceai_reconciliation.reports.report_model import CaseReport


def _table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = str(label)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value if value is not None else "Neidentificat")
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(8)
    return table



def _date_for_operator(value):
    """Display source dates consistently; midnight is not evidence of a time."""
    from datetime import datetime
    raw = str(value or "").strip()
    if not raw or raw == "-":
        return "Neidentificată"
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError:
        return raw
    return (parsed.strftime("%d.%m.%Y %H:%M:%S")
            if parsed.time().isoformat() != "00:00:00"
            else parsed.strftime("%d.%m.%Y"))


def _source_reference(entry):
    """Only expose source identifiers explicitly carried by case evidence."""
    for key in ("source_row", "SourceRowNumber", "source_row_number", "row_number", "source_line"):
        value = entry.get(key)
        if value is not None and str(value).strip() not in ("", "-"):
            return str(value)
    return "Nedisponibil în cronologia cazului"




def _actor_label(entry):
    actor = entry.get("actor") or {}
    user = str(actor.get("user") or "").strip()
    person = str(actor.get("person") or "").strip()
    if user and person:
        return f"{person} / {user}" if user != person else user
    return person or user or "Nedisponibil în sursă"


def _operational_status(*, confirmed_cause: bool = False,
                        operator_approved: bool = False,
                        proposed_correction: dict | None = None) -> str:
    """Fail closed: a computed quantity residual never authorizes a correction."""
    if (operator_approved and proposed_correction
            and all(proposed_correction.get(key) for key in
                    ("system", "document", "operation", "evidence"))):
        return "CORECȚIE PROPUSĂ"
    if confirmed_cause:
        return "EXPLICAT"
    return "REVIEW"


def _decision(doc, checks):
    doc.add_heading("Ce verifică operatorul", level=2)
    for check in checks:
        doc.add_paragraph("☐ " + check)
    doc.add_heading("Decizia operatorului", level=2)
    _table(doc, ["Decizie", "Observații / dovezi suplimentare"], [
        ("☐ Confirmat   ☐ Respins   ☐ Verificare suplimentară", "......................................."),
        ("Verificator / dată", "......................................."),
    ])


def generate_operator_causes_docx(
    report: CaseReport, output_path: str | Path, *,
    details: dict | None = None,
    human_evidence: list[dict] | None = None,
) -> Path:
    """Demo v0.2-like read-only report, with bounded, evidence-grounded pages."""
    details = details or {}
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(28, 72, 106)
    doc.add_heading("RAPORT OPERATOR — CAUSES", 0)
    doc.add_paragraph(f"{report.case_id}  |  {report.report_date}  |  TraceAI Control")
    doc.add_paragraph(f"{report.product_code} — {report.product_name}")
    doc.add_paragraph(f"Perioada: {report.period}   •   Verdict: {report.verdict}   •   Status: {report.case_status}")
    if human_evidence:
        doc.add_heading("Confirmări umane — dovezi asociate sursei", level=1)
        for item in human_evidence:
            doc.add_paragraph(
                f"ID {item['evidence_id']} | {item['system']} rând {item['source_row']} | "
                f"{item['source_date']} | {item['signed_quantity']:+g} {item['unit']}"
            )
            doc.add_paragraph(
                f"Verificare fizică: {'confirmată' if item['physical_count_confirmed'] else 'neconfirmată'}; "
                f"corespondent WME: "
                f"{'nu a fost efectuat' if item['counterpart_wme_exists'] is False else 'de verificat'}."
            )
            doc.add_paragraph(f"Declarație operator: {item['statement']} (sursă: {item['provenance']}).")
        doc.add_paragraph(
            "Confirmarea se aplică doar operațiunii identificate, nu reprezintă MATCH "
            "WMS-WME, nu schimbă verdictul general REVIEW și nu autorizează ajustări."
        )
    doc.add_heading("Rezumat stoc — toate gestiunile", level=1)
    zone_rows = details.get("stock_report_zone_matching", {}).get("rows", [])
    rows = [
        (
            row.get("report_warehouse", "Neidentificat"),
            f'{float(row.get("quantity_wms") or 0):,.3f}',
            f'{float(row.get("quantity_erp") or 0):,.3f}',
            f'{float(row.get("report_difference") or 0):+,.3f}',
        )
        for row in zone_rows
    ]
    if not rows:
        rows = [("Gestiuni indisponibile în datele investigației", "—", "—", "—")]
    _table(doc, ["Zonă / gestiune", "WMS", "WME", "WMS − WME"], rows)
    if zone_rows:
        wms = sum(float(x.get("quantity_wms") or 0) for x in zone_rows)
        wme = sum(float(x.get("quantity_erp") or 0) for x in zone_rows)
        doc.add_paragraph(f"TOTAL ARTICOL: WMS {wms:,.3f} | WME {wme:,.3f} | Diferență {wms-wme:+,.3f} {report.unit}")
    staging = any(f.get("type") == "RECEPTION_STAGING_INCLUDED_IN_STORAGE" for f in report.findings)
    if staging:
        doc.add_paragraph("Regulă confirmată: Recepție nerelocată este zonă operațională a Depozitării; rândurile originale rămân vizibile.")
    mapping_summary = details.get("stock_report_zone_matching") or {}
    explained_by_mapping = (
        bool(zone_rows)
        and mapping_summary.get("unresolved_rows") == 0
        and bool(mapping_summary.get("confirmed_report_mapping_rows"))
        and abs(float(mapping_summary.get("mapping_effect_total") or 0) - (wms - wme)) <= 0.001
    )
    if explained_by_mapping:
        doc.add_paragraph(
            "Diferență în RAPORTARE, explicată prin maparea WME pe gestiuni; "
            "nu este demonstrată o lipsă fizică de stoc. "
            "Validarea tuturor operațiunilor WMS/WME rămâne separată (REVIEW)."
        )
    elif zone_rows and abs(wms-wme) > 0.001:
        doc.add_paragraph(
            "Diferență RAPORTATĂ, cu origine încă de verificat: nu se poate "
            "confirma o diferență fizică sau o eroare operațională din simplul sold. "
            "Nu efectuați ajustări fără validare documentară."
        )
    else:
        doc.add_paragraph("Un sold egal nu dovedește reconcilierea tuturor documentelor.")
    doc.add_heading("Registru CAUSES", 1)
    events = report.operation_timeline
    deliveries = defaultdict(lambda: {"out": [], "ret": [], "ae": []})
    import re
    for entry in events:
        system = str(entry.get("system") or "").upper()
        name = str(entry.get("document") or "").upper().strip()
        operation = str(entry.get("operation_type") or "").upper()
        if system == "WMS":
            m = re.fullmatch(r"SFA(\d+)(-1)?", name)
            if m and operation in ("LIVRARE", "RECEPTIE"):
                key = "ret" if m.group(2) and operation == "RECEPTIE" else "out" if not m.group(2) and operation == "LIVRARE" else None
                if key:
                    deliveries[m.group(1)][key].append(entry)
        elif system == "WME" and operation == "AE" and name.isdigit():
            deliveries[name]["ae"].append(entry)
    flows = []
    for no, x in sorted(deliveries.items()):
        if x["out"] and x["ret"] and x["ae"]:
            net = sum(float(e.get("quantity") or 0) for e in x["out"]+x["ret"])
            erp = sum(float(e.get("quantity") or 0) for e in x["ae"])
            if abs(net-erp) <= .001:
                flows.append((no, x, net))
    # Only document-less stock adjustments, not production consumption postings.
    # Keep opposite-sign entries together: a negative line alone is not a loss.
    wms_corrections = [
        e for e in events
        if str(e.get("system") or "").upper() == "WMS"
        and str(e.get("document") or "").strip() in ("", "-")
        and str(e.get("operation_type") or "").casefold().startswith("ajustare")
    ]
    wms_corrections.sort(key=lambda e: (str(e.get("date") or ""), str(e.get("source_row") or "")))
    correction_groups = defaultdict(list)
    for e in wms_corrections:
        correction_groups[(str(e.get("date") or "")[:10],
                           str((e.get("source_metadata") or {}).get("lot") or "Lot nedisponibil"))].append(e)
    pv = [x for x in report.unmatched_operations if str(x.get("operation_type") or "").upper() in ("PV", "INV-MIN", "INV-PLS")]
    # Warehouse classification comes from an exact WME F source entry,
    # never from the customer's name or a free-text reason.
    wme_f_mose_docs = {
        str(e.get("document") or "").strip().upper()
        for e in events
        if str(e.get("system") or "").upper() == "WME"
        and str(e.get("operation_type") or "").upper() == "F"
        and str(e.get("warehouse") or "").strip().upper() == "MOSE"
    }
    mose = [
        x for x in report.unmatched_operations
        if str(x.get("system") or "").upper() == "WME"
        and str(x.get("operation_type") or "").upper() == "F"
        and str(x.get("document") or "").strip().upper() in wme_f_mose_docs
    ]
    # Multiple investigative views can surface the same economic document.
    # Keep one operator row per exact system, type and document; prefer a dated row.
    def unique_documents(items):
        chosen = {}
        for item in items:
            key = (str(item.get("system") or "").upper(),
                   str(item.get("operation_type") or "").upper(),
                   str(item.get("document") or "").strip().upper())
            current = chosen.get(key)
            if current is None or (str(current.get("date") or "-") == "-"
                                   and str(item.get("date") or "-") != "-"):
                chosen[key] = item
        return list(chosen.values())

    pv = unique_documents(pv)
    mose = unique_documents(mose)
    # Candidate comparison only: the same quantity and a nearby calendar date
    # never constitute a verified inventory reconciliation.
    from datetime import datetime
    pv_candidates = []
    for pv_item in pv:
        if str(pv_item.get("system") or "").upper() != "WME":
            continue
        matching_events = [
            e for e in events
            if str(e.get("system") or "").upper() == "WME"
            and str(e.get("operation_type") or "").upper() == "PV"
            and str(e.get("document") or "").strip() == str(pv_item.get("document") or "").strip()
        ]
        if len(matching_events) != 1:
            continue
        pv_event = matching_events[0]
        try:
            pv_date = datetime.fromisoformat(str(pv_event["date"])[:10])
            pv_quantity = float(pv_event["quantity"])
        except (ValueError, TypeError, KeyError):
            continue
        for e in wms_corrections:
            try:
                correction_date = datetime.fromisoformat(str(e["date"])[:10])
                quantity = float(e["quantity"])
            except (ValueError, TypeError, KeyError):
                continue
            offset_days = (correction_date - pv_date).days
            if abs(offset_days) <= 3 and abs(quantity - pv_quantity) <= 0.001:
                pv_candidates.append((pv_event, e, offset_days))

    timeline_by_document = {}
    for entry in events:
        key = (str(entry.get("system") or "").upper(),
               str(entry.get("document") or "").strip().upper(),
               str(entry.get("operation_type") or "").upper())
        timeline_by_document.setdefault(key, []).append(entry)

    def sourced_row(item):
        system = str(item.get("system") or "").upper()
        document = str(item.get("document") or "").strip().upper()
        operation = str(item.get("operation_type") or "").upper()
        matches = timeline_by_document.get((system, document, operation), [])
        # An inventory without document number cannot be attributed to a
        # timeline row solely on the basis of equal quantities.
        # Synthetic inventory findings do not carry the original WMS row ID.
        # Link for display ONLY when date, signed quantity and direction
        # identify precisely one WMS correction. This is not a WME MATCH.
        if system == "WMS" and operation in ("INV-MIN", "INV-PLS"):
            sign = -1 if operation == "INV-MIN" else 1
            candidates = []
            for e in events:
                if str(e.get("system") or "").upper() != "WMS":
                    continue
                op = str(e.get("operation_type") or "").strip().lower()
                if op != ("ajustare negativa" if sign < 0 else "ajustare pozitiva"):
                    continue
                if str(e.get("date") or "")[:10] != str(item.get("date") or "")[:10]:
                    continue
                try:
                    same_qty = abs(float(e.get("quantity")) - float(item.get("quantity"))) <= 0.001
                except (TypeError, ValueError):
                    same_qty = False
                if same_qty:
                    candidates.append(e)
            if len(candidates) == 1:
                e = candidates[0]
                return (_date_for_operator(e.get("date")), system, document,
                        operation, e.get("quantity"), e.get("warehouse") or "Neidentificată",
                        _source_reference(e), str(item.get("reason") or "-"))
        match = matches[0] if len(matches) == 1 and document not in ("", "-") else None
        # WME invoices can contain several product/lot rows for one F document.
        # The unmatched summary stores an unsigned effect on WMS-WME, not the
        # signed source movement. Reconstruct source-facing values from ALL F
        # rows instead of displaying an absent date and a misleading +quantity.
        if system == "WME" and operation == "F" and document not in ("", "-") and matches:
            mose_rows = [
                e for e in matches
                if str(e.get("warehouse") or "").strip().upper() == "MOSE"
            ]
            if len(mose_rows) == len(matches):
                dates = sorted({_date_for_operator(e.get("date")) for e in matches})
                dates = [d for d in dates if d != "Neidentificată"]
                date_label = " / ".join(dates) if dates else "Neidentificată"
                refs = ", ".join(dict.fromkeys(_source_reference(e) for e in matches))
                quantity = sum(float(e.get("quantity") or 0) for e in matches)
                return (date_label, system, document, operation, quantity,
                        "MOSE", refs, str(item.get("reason") or "-"))
        date = (match.get("date") if match else item.get("date")) or "Neidentificată"
        if date == "-":
            date = "Neidentificată"
        quantity = (match.get("quantity") if match else item.get("quantity"))
        warehouse = match.get("warehouse") if match else "Neidentificată"
        return (_date_for_operator(date), system, document, operation, quantity, warehouse,
                _source_reference(match) if match else _source_reference(item),
                str(item.get("reason") or "-"))

    registru = []
    if staging:
        registru.append(("01", "Recepție nerelocată în Depozitare", "Clarificare operațională confirmată"))
    for no, _, _ in flows[:8]:
        registru.append((f"{len(registru)+1:02}", f"Livrare / retur {no}", "Cantități concordante; verificare documentară"))
    if wms_corrections:
        registru.append((f"{len(registru)+1:02}", "Ajustări WMS fără document — verificare lot și sold net", "REVIEW"))
    if pv:
        registru.append((f"{len(registru)+1:02}", "Inventare și corecții", "REVIEW"))
    if mose:
        registru.append((f"{len(registru)+1:02}", "Documente MOSE", "REVIEW"))
    mapping = details.get("stock_report_zone_matching") or {}
    if mapping.get("confirmed_report_mapping_rows"):
        registru.append((f"{len(registru)+1:02}", "Diferență în raportul de stoc vs. soldul WME pe gestiuni", "Explicație tehnică; REVIEW documentar"))
    if not registru:
        registru.append(("01", "Verificare documentară", "REVIEW"))
    _table(doc, ["ID", "Obiect verificare", "Stadiu"], registru)
    doc.add_paragraph("Fiecare decizie se validează separat. Raportul tehnic și cronologia completă sunt păstrate în JSON.")
    doc.add_paragraph(
        "LEGENDĂ: EXPLICAT = mecanism operațional confirmat, fără a închide articolul; "
        "REVIEW = dovadă sau corespondent încă de verificat; "
        "CORECȚIE PROPUSĂ = numai după confirmarea documentară și aprobarea operatorului. "
        "Niciun status nu autorizează modificări automate WMS/WME."
    )

    counter = 0
    def page(title):
        nonlocal counter
        counter += 1
        doc.add_page_break()
        doc.add_heading(f"CAUSE {counter:02} — {title}", 1)
        doc.add_heading("LOCALIZARE RAPIDĂ ÎN SISTEMELE REALE", 2)
    if mapping.get("confirmed_report_mapping_rows"):
        page("Maparea diferenței raportate pe gestiuni")
        doc.add_paragraph(
            "Stadiu: explicație de mapare identificată de motor; "
            "validarea operațională rămâne REVIEW. Această diferență "
            "poate proveni din raportare, nu reprezintă automat minus fizic."
        )
        _table(doc, ["Gestiune", "WMS raport", "WME raport", "WME fișă magazie",
                     "Diferență raport", "Diferență sisteme"], [
            (
                x.get("report_warehouse", "-"),
                f'{float(x.get("quantity_wms") or 0):,.3f}',
                f'{float(x.get("quantity_erp") or 0):,.3f}',
                f'{float(x.get("wme_balance") or 0):,.3f}',
                f'{float(x.get("report_difference") or 0):+,.3f}',
                f'{float(x.get("actual_system_difference") or 0):+,.3f}',
            )
            for x in mapping.get("rows", [])
        ])
        doc.add_paragraph(
            f"Efect diferență raportată atribuit mapării: "
            f"{float(mapping.get('mapping_effect_total') or 0):+,.3f} {report.unit}. "
            "Aceasta este o explicație pentru discrepanța dintre raport și "
            "fișa de magazie, nu dovada reconcilierii tuturor operațiunilor."
        )
        affected = {
            str(x.get("mapped_wme_warehouse") or "").strip().upper()
            for x in mapping.get("rows", [])
            if str(x.get("status") or "").startswith("REPORT_ERP_ZERO_MAPPING")
        }
        relevant = [
            e for e in events if str(e.get("system") or "").upper() == "WME"
            and str(e.get("warehouse") or "").strip().upper() in affected
        ]
        if relevant:
            doc.add_heading("Operațiuni WME în gestiunea afectată", level=2)
            _table(doc, ["Dată", "Document", "Operațiune", "Cantitate",
                         "Gestiune", "Rând WME", "Utilizator WME"], [
                (_date_for_operator(e.get("date")), e.get("document"),
                 e.get("operation_type"), f'{float(e.get("quantity") or 0):+,.3f}',
                 e.get("warehouse"), _source_reference(e), _actor_label(e))
                for e in relevant[:12]
            ])
        receipt = details.get("receipt_matching") or {}
        delivery = details.get("delivery_matching") or {}
        doc.add_paragraph(
            f"Control documentar agregat: recepții susținute "
            f"{receipt.get('supported_documents', 'necunoscut')}, "
            f"rezidual recepții {receipt.get('residual_total', 'necunoscut')}; "
            f"rezidual livrări {delivery.get('residual_total', 'necunoscut')}. "
            "Verificările agregate nu reprezintă o listă de MATCH pe fiecare rând."
        )
        # Read-only document comparison from original timeline. Document identifiers
        # across WMS and WME may differ; quantity agreement alone is NOT MATCH.
        receipt_wms = [e for e in events if str(e.get("system") or "").upper() == "WMS"
                       and str(e.get("operation_type") or "").casefold() == "receptie"]
        receipt_wme = [e for e in events if str(e.get("system") or "").upper() == "WME"
                       and str(e.get("operation_type") or "").upper() in ("FE", "F")]
        delivery_wms = [e for e in events if str(e.get("system") or "").upper() == "WMS"
                        and str(e.get("operation_type") or "").casefold() == "livrare"]
        delivery_wme = [e for e in events if str(e.get("system") or "").upper() == "WME"
                        and str(e.get("operation_type") or "").upper() == "AE"]
        def grouped_documents(items):
            groups = defaultdict(list)
            for e in items:
                groups[str(e.get("document") or "-")].append(e)
            return [(key, sum(float(e.get("quantity") or 0) for e in rows),
                     ", ".join(sorted({_source_reference(e) for e in rows})[:4]))
                    for key, rows in sorted(groups.items())]
        if receipt_wms or receipt_wme or delivery_wms or delivery_wme:
            doc.add_heading("Documente WMS și WME — control vizual, fără MATCH automat", level=2)
            for label, left, right in (
                ("Recepții", receipt_wms, receipt_wme),
                ("Livrări", delivery_wms, delivery_wme),
            ):
                if not left and not right:
                    continue
                doc.add_paragraph(label + ": documentele sunt grupate separat pe sistem. "
                                  "Un număr sau un total egal nu demonstrează corespondentul.")
                _table(doc, ["Sistem", "Document", "Cantitate", "Rânduri sursă"], [
                    (system, number, f"{qty:+,.3f}", refs)
                    for system, items in (("WMS", left), ("WME", right))
                    for number, qty, refs in grouped_documents(items)
                ][:24])
        
        wme_residual = details.get("wme_residual_movements") or {}
        if wme_residual.get("residual_movements"):
            doc.add_paragraph(
                "ATENȚIE — control de consistență al motorului: "
                "există mișcări WME afișate ca reziduale și operațiuni "
                "inbound considerate concordante. Nu atribuiți acestor "
                "mișcări efectul asupra soldului înainte de eliminarea "
                "eventualei duble clasificări."
            )
            _table(doc, ["Document WME", "Operațiune", "Cantitate WME",
                         "Gestiune", "Clasificare motor"], [
                (x.get("document"), x.get("operation_type"),
                 f'{float(x.get("wme_signed_quantity") or 0):+,.3f}',
                 x.get("warehouse"), x.get("status"))
                for x in wme_residual.get("movements", [])[:10]
            ])
        _decision(doc, [
            "Verifică dacă soldul MOSTRE din fișa WME a fost omis din raport.",
            "Compară NT intern cu mutările WMS, fără a forța MATCH.",
            "Confirmă separat corespondența documentelor de recepție și livrare.",
        ])
    if staging:
        page("Recepție nerelocată în Depozitare")
        doc.add_paragraph("Stadiu operațional: " + _operational_status(confirmed_cause=True))
        _table(doc, ["Gestiune", "WMS", "WME", "Diferență"], rows)
        doc.add_paragraph("Perimetru operațional comun confirmat. Nu înseamnă că toate documentele sunt validate.")
        _decision(doc, ["Confirmă cantitatea rămasă în Recepție.", "Verifică relocarea fără dublare de stoc."])
    for no, match, net in flows[:8]:
        page(f"Livrare {no} și retur parțial")
        doc.add_paragraph("Stadiu operațional: " + _operational_status() + " — retur de validat documentar.")
        entries = match["out"]+match["ret"]+match["ae"]
        _table(doc, ["Sistem", "Data / ora", "Document", "Operațiune", "Cantitate", "Gestiune", "Rând sursă", "Persoană / user"], [
            (e.get("system"), _date_for_operator(e.get("date")),
             e.get("document"), e.get("operation_type"),
             e.get("quantity"), e.get("warehouse"), _source_reference(e), _actor_label(e)) for e in entries
        ])
        metadata_entries = [e for e in match["ret"] if e.get("source_metadata")]
        if metadata_entries:
            doc.add_paragraph("Identificatori expliciți din exportul WMS (retur):")
            _table(doc, ["Document", "Rând sursă", "NIR", "Palet", "Lot", "Locație"], [
                (e.get("document"), _source_reference(e),
                 e["source_metadata"].get("nir", "Nedisponibil"),
                 e["source_metadata"].get("pallet", "Nedisponibil"),
                 e["source_metadata"].get("lot", "Nedisponibil"),
                 e["source_metadata"].get("location", "Nedisponibil"))
                for e in metadata_entries
            ])
        erp = sum(float(e.get("quantity") or 0) for e in match["ae"])
        doc.add_paragraph(f"Calcul: WMS net {net:,.3f} / WME AE {erp:,.3f} / Rezidual {net-erp:+,.3f} {report.unit}.")
        doc.add_paragraph("Egalitatea documentară cantitativă nu confirmă singură că recepția este un retur real.")
        _decision(doc, ["Confirmă returul în documentul de recepție/NIR.", "Verifică produsul, gestiunea, lotul și referințele sursă."])
    if wms_corrections:
        page("Ajustări WMS fără document")
        doc.add_paragraph("Stadiu operațional: REVIEW — corecțiile nu au cauză confirmată.")
        _table(doc, ["Data / ora", "Operațiune", "Cantitate", "Lot", "Rând WMS", "Persoană / user"], [
            (_date_for_operator(e.get("date")), e.get("operation_type"),
             f'{float(e.get("quantity") or 0):+,.3f}',
             (e.get("source_metadata") or {}).get("lot", "Nedisponibil"),
             _source_reference(e), _actor_label(e))
            for e in wms_corrections[:12]
        ])
        detail_rows = [
            (_source_reference(e), (e.get("wms_details") if str(e.get("wms_details") or "").strip() else "motiv necompletat"))
            for e in wms_corrections[:12]
        ]
        doc.add_paragraph("Detalii WMS — observație introdusă manual de o persoană în WMS, păstrată textual din export. Textul indică ce a consemnat operatorul, dar nu confirmă singur executarea ori corespondentul WME:")
        _table(doc, ["Rând WMS", "Detalii"], detail_rows)
        for (day, lot), items in sorted(correction_groups.items()):
            total = sum(float(e.get("quantity") or 0) for e in items)
            doc.add_paragraph(
                f"Grup {day} / lot {lot}: {len(items)} mișcări, "
                f"net {total:+,.3f} {report.unit}. "
                "Efect net calculat, nu cauză confirmată."
            )
        if len(wms_corrections) > 12:
            doc.add_paragraph(
                f"Alte {len(wms_corrections)-12} ajustări sunt păstrate în JSON; "
                "lista din DOCX este limitată pentru lizibilitate."
            )
        doc.add_paragraph(
            "O coincidență între netul unei zile și diferența finală este doar "
            "un candidat documentar. Perechile plus/minus se evaluează împreună."
        )
        _decision(doc, [
            "Identifică documentul justificativ pentru fiecare ajustare fără număr.",
            "Verifică lotul, data, semnul și eventualele perechi plus/minus.",
            "Compară PV WME și ajustările WMS fără a presupune corespondentul."
        ])
    if pv_candidates:
        doc.add_paragraph("CORESPONDENȚE PV CANDIDATE — NU SUNT MATCH CONFIRMATE")
        _table(doc, ["PV WME", "Rând WME", "Ajustare WMS", "Rând WMS", "Decalaj zile", "Cantitate"], [
            (pv_event.get("document"), _source_reference(pv_event),
             _date_for_operator(wms_event.get("date")), _source_reference(wms_event),
             str(offset_days), f'{float(wms_event.get("quantity") or 0):+,.3f}')
            for pv_event, wms_event, offset_days in pv_candidates[:12]
        ])
        doc.add_paragraph(
            "Aceeași cantitate și o dată apropiată nu demonstrează legătura dintre documente. "
            "Operatorul verifică PV-ul, lotul, gestiunea, identificatorul operațiunii și aprobarea."
        )
    for title, items, checks in [
        ("Inventare și corecții",pv,["Găsește PV-ul și ajustarea WMS.", "Verifică data, cantitatea, semnul, lotul și gestiunea."]),
        ("Documente F în MOSE",mose,["Verifică factura F și clientul înscris pe factură.", "Caută NT WME de intrare în gestiunea MOSE și livrarea WMS TRANSFER_MOSE corespunzătoare.", "Nu confunda gestiunea MOSE cu clientul MOSE Business Distribution."]),
    ]:
        if not items:
            continue
        page(title)
        doc.add_paragraph("Stadiu operațional: " + _operational_status())
        _table(doc, ["Data/ora sursă", "Sistem", "Document", "Operațiune", "Cantitate mișcare", "Gestiune", "Rând sursă", "Motiv / efect tehnic"], [
            sourced_row(x) for x in items[:15]
        ])
        doc.add_paragraph(
            "Cantitate mișcare = semnul din cronologia sursă când există corespondență exactă. "
            "Efectul tehnic asupra diferenței, dacă apare în motiv, este alt indicator și "
            "nu trebuie confundat cu semnul mișcării."
        )
        if title == "Documente F în MOSE":
            doc.add_paragraph(
                "REGULĂ MOSE CONFIRMATĂ: MOSE este gestiune exclusiv WME; "
                "nu există gestiune MOSE și nici sold MOSE în WMS. "
                "Marfa iese din Depozit Principal prin Livrare WMS TRANSFER_MOSE "
                "și intră în MOSE prin NT WME. Factura F este o ieșire ulterioară "
                "din gestiunea MOSE către clientul înscris pe document; "
                "nu reprezintă transferul în sine. "
                "MOSE Business Distribution este un client separat, nu gestiunea MOSE."
            )
            transfer_info = details.get("transfer_mose_matching") or {}
            matched = transfer_info.get("matches")
            if matched is not None:
                doc.add_paragraph(
                    f"Transferuri interne WMS ↔ NT WME identificate de motor: {matched}. "
                    "Această asociere nu validează automat facturile F."
                )
        # Source-user attribution is already recorded in JSON; avoid
        # duplicating PV/F economic documents in a second DOCX table.
        # Only emit a short attribution note when a source actor exists.
        actor_labels = []
        for item in items[:15]:
            key = (str(item.get("system") or "").upper(),
                   str(item.get("document") or "").strip().upper(),
                   str(item.get("operation_type") or "").upper())
            matches = timeline_by_document.get(key, [])
            if len(matches) == 1 and key[1] not in ("", "-"):
                if matches[0].get("actor"):
                    actor_labels.append(f"{key[1]}: {_actor_label(matches[0])}")
        if actor_labels:
            doc.add_paragraph("Persoană / user (sursă): " + "; ".join(actor_labels))
        doc.add_paragraph("REVIEW: documentele nu sunt considerate cauze confirmate.")
        _decision(doc, checks)
    doc.add_paragraph("NOTĂ: raportul nu autorizează ajustări WMS/WME. Ora 00:00 din WME poate însemna doar dată calendaristică.")
    doc.save(target)
    return target
