from __future__ import annotations

import csv
import hashlib
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence
from zipfile import ZipFile

import pandas as pd
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


ALIASES: dict[str, tuple[str, ...]] = {
    "product_code": (
        "CodArticol", "Cod articol", "Cod produs", "Cod DS", "Cod WMS",
        "Cod produs WMS/DS", "DS", "ProductCode",
    ),
    "product_name": (
        "DenumireArticol", "Denumire articol", "Denumire produs", "Produs",
        "ProductName", "productname", "Denumire",
    ),
    "quantity": (
        "Cantitate", "cantitate", "Quantity", "Qty", "Cant", "Cantitate operata",
        "SignedQuantity",
    ),
    "unit": ("UM", "U.M.", "Unitate masura", "Unit", "U.M."),
    "date": (
        "Data", "Data operatie", "Data operatiune", "Data document", "Date",
        "DataDoc", "Data Start",
    ),
    "document": (
        "Document", "NrDocument", "Nr document", "Numar document", "Numar",
        "NrDoc", "Comanda", "Numar Comanda", "Numar comanda", "Order",
    ),
    "operation_type": (
        "Tip operatiune", "Tip operatie", "Operatiune", "Tip document",
        "TipDoc", "Tip", "OperationType", "TipLinie",
    ),
    "warehouse": (
        "Gestiune", "Depozit", "Locatie", "Warehouse", "Gestiune sursa",
    ),
    "running_balance": (
        "Stoc", "Sold", "Stoc la data", "Sold final", "Balance",
    ),
    "wme_item_id": (
        "ID articol WME", "Id articol WME", "ID articol", "IdArticol",
        "idArticol", "Cod articol WME", "CodArticolWME", "ArticolID",
        "WMEItemId",
    ),
    "wms_product_code": (
        "Cod produs WMS/DS", "Cod WMS/DS", "Cod DS", "Cod WMS",
        "CodArticolWMS", "CodArticol", "Cod produs", "Cod articol",
    ),
    "stock_wme": ("QuantityERP", "Stoc ERP", "Stoc WME", "QuantityWME"),
    "stock_wms": ("QuantityWMS", "Stoc WMS"),
    "stock_difference": ("QuantityDifference", "Diferenta", "Diferență"),
    "initial_stock": (
        "StocInitial", "Stoc initial", "Stoc inițial", "Quantity",
        "Cantitate", "cantitate", "Stoc",
    ),
    "lot": ("Lot", "Lot produs", "LotArticol"),
}


def normalize_header(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().strip()
    return re.sub(r"[^a-z0-9]+", "", text)


def resolve_column(
    columns: Sequence[object], semantic_name: str, required: bool = True
) -> str | None:
    candidates = ALIASES.get(semantic_name, ())
    normalized = {normalize_header(col): str(col) for col in columns}
    for alias in candidates:
        key = normalize_header(alias)
        if key in normalized:
            return normalized[key]
    if required:
        raise ValueError(
            f"Nu pot rezolva coloana '{semantic_name}'. "
            f"Coloane disponibile: {list(map(str, columns))}"
        )
    return None


def sha256_file(path: str | Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        while True:
            chunk = handle.read(chunk_size)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def detect_csv_delimiter(path: str | Path) -> str:
    path = Path(path)
    sample = path.read_bytes()[:65536]
    for encoding in ("utf-8-sig", "utf-8", "cp1250", "latin1"):
        try:
            text = sample.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        return ";"

    first_line = text.splitlines()[0] if text.splitlines() else text
    counts = {delimiter: first_line.count(delimiter) for delimiter in (";", ",", "\t", "|")}
    best = max(counts, key=counts.get)
    if counts[best] > 0:
        return best

    try:
        return csv.Sniffer().sniff(text, delimiters=",;\t|").delimiter
    except csv.Error:
        return ";"


def parse_period(period: str) -> tuple[pd.Timestamp, pd.Timestamp]:
    tokens = re.findall(r"\d{2}\.\d{2}\.\d{4}|\d{4}-\d{2}-\d{2}", str(period))
    if len(tokens) != 2:
        raise ValueError(
            "Perioada trebuie să conțină exact două date, de exemplu "
            "'01.01.2026-04.10.2026'."
        )

    def parse_one(value: str) -> pd.Timestamp:
        if "." in value:
            return pd.Timestamp(datetime.strptime(value, "%d.%m.%Y"))
        return pd.Timestamp(datetime.strptime(value, "%Y-%m-%d"))

    start, end = map(parse_one, tokens)
    if start > end:
        raise ValueError("Data de început este după data de sfârșit.")
    return start, end


def parse_operational_dates(series: pd.Series) -> pd.Series:
    """Parse operational dates using explicit source conventions.

    Supported formats:
    - WMS / production exports: month-first slash timestamps with AM/PM;
    - WME / Romanian exports: day-first dotted dates;
    - ISO timestamps;
    - Excel serial dates in the plausible operational range.

    Ambiguous slash dates without AM/PM are parsed month-first because the
    current WMS/production exports use US-style month/day ordering.
    """
    parsed = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")

    numeric = pd.to_numeric(series, errors="coerce")
    excel_serial = numeric.between(20000, 80000, inclusive="both")
    if excel_serial.any():
        parsed.loc[excel_serial] = pd.to_datetime(
            numeric.loc[excel_serial],
            unit="D",
            origin="1899-12-30",
            errors="coerce",
        )

    remaining = ~excel_serial
    if not remaining.any():
        return parsed

    text = series.loc[remaining]
    text_string = text.astype(str).str.strip()

    dotted = text_string.str.match(
        r"^\d{1,2}\.\d{1,2}\.\d{4}(?:\s|$)",
        na=False,
    )
    if dotted.any():
        parsed.loc[text.index[dotted]] = pd.to_datetime(
            text.loc[dotted],
            errors="coerce",
            dayfirst=True,
        )

    slash = text_string.str.match(
        r"^\d{1,2}/\d{1,2}/\d{4}(?:\s|$)",
        na=False,
    )
    if slash.any():
        parsed.loc[text.index[slash]] = pd.to_datetime(
            text.loc[slash],
            errors="coerce",
            dayfirst=False,
        )

    iso = text_string.str.match(r"^\d{4}-\d{2}-\d{2}", na=False)
    if iso.any():
        parsed.loc[text.index[iso]] = pd.to_datetime(
            text.loc[iso],
            errors="coerce",
            yearfirst=True,
        )

    handled = dotted | slash | iso
    fallback = ~handled
    if fallback.any():
        parsed.loc[text.index[fallback]] = pd.to_datetime(
            text.loc[fallback],
            errors="coerce",
            format="mixed",
            dayfirst=True,
        )

    return parsed


def filter_by_period(
    df: pd.DataFrame,
    period: str,
    semantic_date_column: str = "date",
) -> pd.DataFrame:
    if df.empty:
        return df.copy()
    date_col = resolve_column(df.columns, semantic_date_column, required=False)
    if date_col is None:
        return df.copy()

    start, end = parse_period(period)
    parsed = parse_operational_dates(df[date_col])
    mask = parsed.between(start, end + pd.Timedelta(days=1) - pd.Timedelta(microseconds=1))
    result = df.loc[mask].copy()
    result[date_col] = parsed.loc[mask]
    return result


def _string_mask(series: pd.Series, accepted_values: set[str]) -> pd.Series:
    normalized = series.astype(str).str.strip()
    return normalized.isin(accepted_values)


def filter_csv_by_values(
    path: str | Path,
    semantic_column: str,
    accepted_values: Iterable[str],
    chunksize: int = 50000,
) -> pd.DataFrame:
    path = Path(path)
    accepted = {str(value).strip() for value in accepted_values}
    delimiter = detect_csv_delimiter(path)
    chunks: list[pd.DataFrame] = []
    resolved_column: str | None = None

    for encoding in ("utf-8-sig", "cp1250", "latin1"):
        try:
            iterator = pd.read_csv(
                path,
                sep=delimiter,
                chunksize=chunksize,
                low_memory=False,
                encoding=encoding,
            )
            for chunk in iterator:
                if resolved_column is None:
                    resolved_column = resolve_column(chunk.columns, semantic_column)
                mask = _string_mask(chunk[resolved_column], accepted)
                if mask.any():
                    chunks.append(chunk.loc[mask].copy())
            break
        except UnicodeDecodeError:
            chunks = []
            resolved_column = None
            continue

    if not chunks:
        return pd.DataFrame()
    return pd.concat(chunks, ignore_index=True)


def read_excel_table(path: str | Path, sheet_name: str | int = 0) -> pd.DataFrame:
    return pd.read_excel(Path(path), sheet_name=sheet_name)


def filter_xlsx_by_values(
    path: str | Path,
    semantic_column: str,
    accepted_values: Iterable[str],
) -> pd.DataFrame:
    path = Path(path)
    accepted = {str(value).strip() for value in accepted_values}
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)

    try:
        headers = list(next(rows))
    except StopIteration:
        workbook.close()
        return pd.DataFrame()

    column = resolve_column(headers, semantic_column)
    index = headers.index(column)
    matches: list[tuple[object, ...]] = []

    for row in rows:
        value = row[index] if index < len(row) else None
        if str(value).strip() in accepted:
            matches.append(tuple(row))

    workbook.close()
    if not matches:
        return pd.DataFrame(columns=headers)
    return pd.DataFrame(matches, columns=headers)


def read_filtered(
    path: str | Path,
    semantic_column: str,
    accepted_values: Iterable[str],
) -> pd.DataFrame:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return filter_csv_by_values(path, semantic_column, accepted_values)
    if suffix == ".xlsx":
        return filter_xlsx_by_values(path, semantic_column, accepted_values)
    if suffix == ".xls":
        raise ValueError(
            f"{path.name}: format .xls neacceptat în citirea filtrată. "
            "Convertește sursa la .xlsx fără a modifica datele."
        )
    raise ValueError(f"Format sursă neacceptat: {path.suffix}")


def read_production_product(
    path: str | Path,
    product_code: str,
    period: str | None = None,
    chunksize: int = 50000,
) -> pd.DataFrame:
    path = Path(path)
    delimiter = detect_csv_delimiter(path)
    matches: list[pd.DataFrame] = []

    iterator = pd.read_csv(
        path,
        sep=delimiter,
        chunksize=chunksize,
        low_memory=False,
        encoding="utf-8-sig",
    )
    for chunk in iterator:
        direct = resolve_column(chunk.columns, "product_code", required=False)
        if direct is not None:
            mask = chunk[direct].astype(str).str.strip().eq(str(product_code).strip())
        else:
            pre = next(
                (c for c in chunk.columns if normalize_header(c) == normalize_header("PRE_Cod Articol")),
                None,
            )
            con = next(
                (c for c in chunk.columns if normalize_header(c) == normalize_header("CON_Cod Articol")),
                None,
            )
            if pre is None and con is None:
                raise ValueError(
                    f"{path.name}: nu găsesc PRE_Cod Articol / CON_Cod Articol."
                )
            mask = pd.Series(False, index=chunk.index)
            if pre is not None:
                mask |= chunk[pre].astype(str).str.strip().eq(str(product_code).strip())
            if con is not None:
                mask |= chunk[con].astype(str).str.strip().eq(str(product_code).strip())

        if mask.any():
            matches.append(chunk.loc[mask].copy())

    result = pd.concat(matches, ignore_index=True) if matches else pd.DataFrame()
    return filter_by_period(result, period) if period and not result.empty else result


WME_STANDARD_COLUMNS = [
    "LineNo",
    "OperationType",
    "Document",
    "Date",
    "Incoming",
    "Outgoing",
    "Balance",
    "Price",
    "Account",
    "Partner",
    "ProductName",
    "CounterWarehouse",
    "WMEItemId",
    "Warehouse",
    "Unit",
    "Note",
    "InternalRef",
    "SignedQuantity",
]


def read_wme_product(
    path: str | Path,
    product_name: str,
    period: str | None = None,
) -> tuple[pd.DataFrame, set[str]]:
    """
    Read one WME product from a large XLSX export.

    Fast path: when ProductName exists in sharedStrings.xml, read the worksheet
    XML once, locate only the shared-string index used in column K, and decode
    only those matching rows. This avoids Python-level iteration across ~243k
    WME rows for a single-case self-service run.

    Conservative fallback: if the requested product is not represented by a
    shared-string index, use streaming row iteration without changing matching
    semantics.
    """
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError(
            f"{path.name}: fișa WME trebuie să fie .xlsx pentru citirea targetată."
        )

    xml_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    namespace_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    column_pattern = re.compile(r"([A-Z]+)")
    accepted_columns = set("ABCDEFGHIJKLMNOPQ")
    target = str(product_name).strip()

    def decode_cell(
        cell: etree._Element,
        shared_strings: list[str],
    ) -> object:
        cell_type = cell.get("t")
        value_node = cell.find(xml_ns + "v")
        value: object = value_node.text if value_node is not None else None
        if cell_type == "s" and value is not None:
            try:
                value = shared_strings[int(value)]
            except (ValueError, IndexError):
                return value
        elif cell_type == "inlineStr":
            inline = cell.find(xml_ns + "is")
            if inline is not None:
                value = "".join(
                    (node.text or "")
                    for node in inline.iter(xml_ns + "t")
                )
        return value

    def decoded_row(
        row_element: etree._Element,
        shared_strings: list[str],
    ) -> dict[str, object]:
        values: dict[str, object] = {}
        for cell in row_element.findall(xml_ns + "c"):
            reference = cell.get("r") or ""
            match = column_pattern.match(reference)
            if match is None:
                continue
            column = match.group(1)
            if column in accepted_columns:
                values[column] = decode_cell(cell, shared_strings)

        def numeric(column: str) -> float:
            value = values.get(column)
            try:
                return float(value)
            except (TypeError, ValueError):
                return 0.0

        raw_date = values.get("D")
        parsed_date: object = raw_date
        if raw_date not in (None, ""):
            try:
                parsed_date = from_excel(float(raw_date))
            except (TypeError, ValueError, OverflowError):
                parsed_date = raw_date

        incoming = numeric("E")
        outgoing = numeric("F")
        item_id = str(values.get("M") or "").strip()
        return {
            "LineNo": values.get("A"),
            "OperationType": values.get("B"),
            "Document": values.get("C"),
            "Date": parsed_date,
            "Incoming": incoming,
            "Outgoing": outgoing,
            "Balance": numeric("G"),
            "Price": numeric("H"),
            "Account": values.get("I"),
            "Partner": values.get("J"),
            "ProductName": str(values.get("K") or "").strip(),
            "CounterWarehouse": values.get("L"),
            "WMEItemId": item_id,
            "Warehouse": values.get("N"),
            "Unit": values.get("O"),
            "Note": values.get("P"),
            "InternalRef": values.get("Q"),
            "SignedQuantity": incoming - outgoing,
        }

    with ZipFile(path) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = etree.fromstring(archive.read("xl/sharedStrings.xml"))
            shared_strings = [
                "".join(
                    (node.text or "")
                    for node in item.iter(xml_ns + "t")
                )
                for item in root.findall(xml_ns + "si")
            ]

        target_shared_indexes = {
            index
            for index, value in enumerate(shared_strings)
            if str(value).strip() == target
        }
        worksheet_names = sorted(
            name
            for name in archive.namelist()
            if name.startswith("xl/worksheets/sheet")
            and name.endswith(".xml")
        )
        if not worksheet_names:
            raise ValueError(
                f"{path.name}: nu găsesc worksheet XML pentru fișa WME."
            )

        rows: list[dict[str, object]] = []
        ids: set[str] = set()

        if target_shared_indexes:
            worksheet = archive.read(worksheet_names[0])
            row_ranges: set[tuple[int, int]] = set()

            for shared_index in target_shared_indexes:
                needle = f"<v>{shared_index}</v>".encode("ascii")
                cursor = 0
                while True:
                    hit = worksheet.find(needle, cursor)
                    if hit < 0:
                        break
                    cursor = hit + len(needle)

                    cell_start = worksheet.rfind(
                        b"<c",
                        max(0, hit - 2048),
                        hit,
                    )
                    cell_end = worksheet.find(
                        b"</c>",
                        hit,
                        min(len(worksheet), hit + 4096),
                    )
                    if cell_start < 0 or cell_end < 0:
                        continue

                    cell_xml = worksheet[cell_start : cell_end + 4]
                    if b't="s"' not in cell_xml:
                        continue
                    reference_match = re.search(
                        rb'\br="K\d+"',
                        cell_xml,
                    )
                    if reference_match is None:
                        continue

                    row_start = worksheet.rfind(
                        b"<row",
                        max(0, cell_start - 65536),
                        cell_start,
                    )
                    row_end = worksheet.find(
                        b"</row>",
                        cell_end,
                        min(len(worksheet), cell_end + 65536),
                    )
                    if row_start < 0 or row_end < 0:
                        continue
                    row_ranges.add((row_start, row_end + 6))

            wrapper_start = (
                f'<root xmlns="{namespace_uri}">'.encode("utf-8")
            )
            wrapper_end = b"</root>"
            for row_start, row_end in sorted(row_ranges):
                fragment = (
                    wrapper_start
                    + worksheet[row_start:row_end]
                    + wrapper_end
                )
                root = etree.fromstring(fragment)
                if len(root) != 1:
                    continue
                item = decoded_row(root[0], shared_strings)
                if item["ProductName"] != target:
                    continue
                item_id = str(item["WMEItemId"] or "").strip()
                if item_id:
                    ids.add(item_id)
                rows.append(item)
        else:
            with archive.open(worksheet_names[0]) as handle:
                context = etree.iterparse(
                    handle,
                    events=("end",),
                    tag=xml_ns + "row",
                )
                for _, row_element in context:
                    product_matches = False
                    for cell in row_element.findall(xml_ns + "c"):
                        reference = cell.get("r") or ""
                        match = column_pattern.match(reference)
                        if match is None or match.group(1) != "K":
                            continue
                        value = decode_cell(cell, shared_strings)
                        product_matches = (
                            str(value or "").strip() == target
                        )
                        break

                    if product_matches:
                        item = decoded_row(row_element, shared_strings)
                        item_id = str(item["WMEItemId"] or "").strip()
                        if item_id:
                            ids.add(item_id)
                        rows.append(item)

                    row_element.clear()
                    parent = row_element.getparent()
                    if parent is not None:
                        while row_element.getprevious() is not None:
                            del parent[0]

    result = pd.DataFrame(rows, columns=WME_STANDARD_COLUMNS)
    if period and not result.empty:
        result = filter_by_period(result, period)
    return result, ids


def numeric_sum(df: pd.DataFrame, semantic_column: str) -> float:
    if df.empty:
        return 0.0
    column = resolve_column(df.columns, semantic_column)
    values = pd.to_numeric(df[column], errors="coerce").fillna(0)
    return float(values.sum())


def first_non_empty(df: pd.DataFrame, semantic_column: str) -> str | None:
    if df.empty:
        return None
    column = resolve_column(df.columns, semantic_column, required=False)
    if column is None:
        return None
    values = df[column].dropna().astype(str).str.strip()
    values = values[values != ""]
    return values.iloc[0] if not values.empty else None
