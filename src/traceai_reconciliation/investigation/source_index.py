from __future__ import annotations

import os
import sqlite3
from datetime import date, datetime
from pathlib import Path
from typing import Iterator

import pandas as pd
from openpyxl import load_workbook

from .source_readers import (
    WME_STANDARD_COLUMNS,
    detect_csv_delimiter,
    filter_by_period,
)


def _fingerprint(path: str | Path) -> dict[str, int | str]:
    source = Path(path).resolve()
    stat = source.stat()
    return {
        "path": str(source),
        "size": int(stat.st_size),
        "mtime_ns": int(stat.st_mtime_ns),
    }


def _sqlite_value(value: object) -> object:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ")
    if isinstance(value, (str, int, float, bytes)):
        return value
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return str(value)


def _csv_chunks(
    path: Path,
    chunksize: int = 50000,
) -> Iterator[pd.DataFrame]:
    delimiter = detect_csv_delimiter(path)
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp1250", "latin1"):
        try:
            iterator = pd.read_csv(
                path,
                sep=delimiter,
                chunksize=chunksize,
                low_memory=False,
                encoding=encoding,
            )
            for chunk in iterator:
                yield chunk
            return
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
    if last_error:
        raise last_error


class SQLiteSourceIndex:
    """
    Local reusable index for the three large operational sources.

    The index is rebuilt atomically whenever path/size/mtime_ns changes for
    WMS, WME or production. RAW files are never modified.
    """

    def __init__(self, cache_dir: str | Path) -> None:
        self.cache_dir = Path(cache_dir)
        self.db_path = self.cache_dir / "operational_sources.sqlite"

    def _expected_meta(
        self,
        wms_path: str | Path,
        wme_path: str | Path,
        production_path: str | Path,
    ) -> dict[str, dict[str, int | str]]:
        return {
            "wms": _fingerprint(wms_path),
            "wme": _fingerprint(wme_path),
            "production": _fingerprint(production_path),
        }

    def is_current(
        self,
        wms_path: str | Path,
        wme_path: str | Path,
        production_path: str | Path,
    ) -> bool:
        if not self.db_path.exists():
            return False

        expected = self._expected_meta(
            wms_path,
            wme_path,
            production_path,
        )
        try:
            with sqlite3.connect(self.db_path) as connection:
                rows = connection.execute(
                    """
                    SELECT source_key, path, size, mtime_ns
                    FROM _source_meta
                    """
                ).fetchall()
        except sqlite3.Error:
            return False

        current = {
            str(source_key): {
                "path": str(path),
                "size": int(size),
                "mtime_ns": int(mtime_ns),
            }
            for source_key, path, size, mtime_ns in rows
        }
        return current == expected

    def ensure_index(
        self,
        wms_path: str | Path,
        wme_path: str | Path,
        production_path: str | Path,
    ) -> bool:
        """
        Ensure a valid index exists.

        Returns True when a rebuild was performed, False when the existing
        index was reused.
        """
        if self.is_current(wms_path, wme_path, production_path):
            return False

        self.cache_dir.mkdir(parents=True, exist_ok=True)
        temporary = self.db_path.with_suffix(".sqlite.tmp")
        temporary.unlink(missing_ok=True)

        try:
            self._build(
                temporary,
                Path(wms_path),
                Path(wme_path),
                Path(production_path),
            )
            os.replace(temporary, self.db_path)
        finally:
            temporary.unlink(missing_ok=True)
        return True

    def _build(
        self,
        target: Path,
        wms_path: Path,
        wme_path: Path,
        production_path: Path,
    ) -> None:
        with sqlite3.connect(target) as connection:
            connection.execute("PRAGMA journal_mode=OFF")
            connection.execute("PRAGMA synchronous=OFF")
            connection.execute("PRAGMA temp_store=MEMORY")

            first = True
            for chunk in _csv_chunks(wms_path):
                chunk.to_sql(
                    "wms",
                    connection,
                    if_exists="replace" if first else "append",
                    index=False,
                )
                first = False
            if first:
                raise ValueError(f"Sursa WMS este goală: {wms_path}")

            first = True
            for chunk in _csv_chunks(production_path):
                chunk.to_sql(
                    "production",
                    connection,
                    if_exists="replace" if first else "append",
                    index=False,
                )
                first = False
            if first:
                raise ValueError(
                    f"Sursa de producție este goală: {production_path}"
                )

            self._build_wme(connection, wme_path)

            wms_columns = {
                row[1]
                for row in connection.execute(
                    "PRAGMA table_info(wms)"
                ).fetchall()
            }
            if "Cod articol" not in wms_columns:
                raise ValueError(
                    "Index WMS: coloana 'Cod articol' lipsește."
                )

            production_columns = {
                row[1]
                for row in connection.execute(
                    "PRAGMA table_info(production)"
                ).fetchall()
            }
            required_production = {
                "PRE_Cod Articol",
                "CON_Cod Articol",
            }
            missing_production = (
                required_production - production_columns
            )
            if missing_production:
                raise ValueError(
                    "Index producție: coloane lipsă "
                    f"{sorted(missing_production)}"
                )

            connection.execute(
                'CREATE INDEX idx_wms_product '
                'ON wms("Cod articol")'
            )
            connection.execute(
                'CREATE INDEX idx_prod_pre '
                'ON production("PRE_Cod Articol")'
            )
            connection.execute(
                'CREATE INDEX idx_prod_con '
                'ON production("CON_Cod Articol")'
            )
            connection.execute(
                "CREATE INDEX idx_wme_product "
                "ON wme(ProductName)"
            )

            connection.execute(
                """
                CREATE TABLE _source_meta (
                    source_key TEXT PRIMARY KEY,
                    path TEXT NOT NULL,
                    size INTEGER NOT NULL,
                    mtime_ns INTEGER NOT NULL
                )
                """
            )
            meta = self._expected_meta(
                wms_path,
                wme_path,
                production_path,
            )
            connection.executemany(
                """
                INSERT INTO _source_meta
                (source_key, path, size, mtime_ns)
                VALUES (?, ?, ?, ?)
                """,
                [
                    (
                        key,
                        value["path"],
                        value["size"],
                        value["mtime_ns"],
                    )
                    for key, value in meta.items()
                ],
            )
            connection.commit()

    def _build_wme(
        self,
        connection: sqlite3.Connection,
        path: Path,
        batch_size: int = 5000,
    ) -> None:
        connection.execute(
            """
            CREATE TABLE wme (
                LineNo,
                OperationType,
                Document,
                Date,
                Incoming REAL,
                Outgoing REAL,
                Balance REAL,
                Price REAL,
                Account,
                Partner,
                ProductName,
                CounterWarehouse,
                WMEItemId,
                Warehouse,
                Unit,
                Note,
                InternalRef,
                SignedQuantity REAL
            )
            """
        )

        placeholders = ",".join("?" for _ in WME_STANDARD_COLUMNS)
        insert_sql = (
            "INSERT INTO wme ("
            + ",".join(WME_STANDARD_COLUMNS)
            + ") VALUES ("
            + placeholders
            + ")"
        )

        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )
        worksheet = workbook[workbook.sheetnames[0]]
        batch: list[tuple[object, ...]] = []

        try:
            for row in worksheet.iter_rows(values_only=True):
                if len(row) < 15:
                    continue

                product_name = str(row[10] or "").strip()
                item_id = str(row[12] or "").strip()
                if not product_name or not item_id:
                    continue

                incoming = pd.to_numeric(
                    pd.Series([row[4]]),
                    errors="coerce",
                ).iloc[0]
                outgoing = pd.to_numeric(
                    pd.Series([row[5]]),
                    errors="coerce",
                ).iloc[0]
                incoming = (
                    0.0 if pd.isna(incoming) else float(incoming)
                )
                outgoing = (
                    0.0 if pd.isna(outgoing) else float(outgoing)
                )

                values = (
                    row[0],
                    row[1],
                    row[2],
                    row[3],
                    incoming,
                    outgoing,
                    row[6],
                    row[7],
                    row[8],
                    row[9],
                    product_name,
                    row[11],
                    item_id,
                    row[13],
                    row[14],
                    row[15] if len(row) > 15 else None,
                    row[16] if len(row) > 16 else None,
                    incoming - outgoing,
                )
                batch.append(
                    tuple(_sqlite_value(value) for value in values)
                )

                if len(batch) >= batch_size:
                    connection.executemany(insert_sql, batch)
                    batch.clear()

            if batch:
                connection.executemany(insert_sql, batch)
        finally:
            workbook.close()

    def read_wms_product(
        self,
        product_code: str,
        period: str | None = None,
    ) -> pd.DataFrame:
        with sqlite3.connect(self.db_path) as connection:
            frame = pd.read_sql_query(
                """
                SELECT *
                FROM wms
                WHERE TRIM(CAST("Cod articol" AS TEXT)) = ?
                """,
                connection,
                params=(str(product_code).strip(),),
            )
        if period and not frame.empty:
            frame = filter_by_period(frame, period)
        return frame

    def read_wme_product(
        self,
        product_name: str,
        period: str | None = None,
    ) -> tuple[pd.DataFrame, set[str]]:
        with sqlite3.connect(self.db_path) as connection:
            frame = pd.read_sql_query(
                """
                SELECT *
                FROM wme
                WHERE ProductName = ?
                """,
                connection,
                params=(str(product_name).strip(),),
            )

        ids = {
            str(value).strip()
            for value in frame.get(
                "WMEItemId",
                pd.Series(dtype=object),
            ).dropna()
            if str(value).strip()
        }
        if period and not frame.empty:
            frame = filter_by_period(frame, period)
        return frame, ids

    def read_production_product(
        self,
        product_code: str,
        period: str | None = None,
    ) -> pd.DataFrame:
        code = str(product_code).strip()
        with sqlite3.connect(self.db_path) as connection:
            frame = pd.read_sql_query(
                """
                SELECT *
                FROM production
                WHERE TRIM(CAST("PRE_Cod Articol" AS TEXT)) = ?
                   OR TRIM(CAST("CON_Cod Articol" AS TEXT)) = ?
                """,
                connection,
                params=(code, code),
            )
        if period and not frame.empty:
            frame = filter_by_period(frame, period)
        return frame
